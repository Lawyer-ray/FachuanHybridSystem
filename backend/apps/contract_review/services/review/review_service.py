from __future__ import annotations

import logging
import uuid
from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from io import BytesIO
from pathlib import Path

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import UploadedFile
from django.db.models import Max
from docx import Document
from docx.document import Document as DocumentType

from apps.contract_review.models.review_revision import ReviewRevision, RevisionKind, RevisionSeverity, RevisionStatus
from apps.contract_review.models.review_task import ProcessStep, ReviewTask, TaskStatus
from apps.contract_review.repositories.review_task_repository import ReviewTaskRepository
from apps.contract_review.services.extraction.title_extractor import generate_output_filename
from apps.contract_review.services.formatting.docx_revision_tool import enable_track_changes
from apps.core.exceptions import ValidationException
from apps.core.filesystem.upload_paths import MediaEntity
from apps.core.llm.service import LLMService, get_llm_service
from apps.core.services.storage_service import sanitize_upload_filename, to_media_abs

from ..exceptions import ContractReviewError, ExtractionError
from ..extraction.content_extractor import ContentExtractor
from ..extraction.heading_numbering import HeadingNumbering
from ..extraction.title_extractor import TitleExtractor
from ..formatting.docx_formatter import DocxFormatter
from ..formatting.docx_revision_tool import DocxRevisionTool
from ..formatting.page_numbering import PageNumbering
from .contract_reviewer import ContractReviewer, ReviewResult
from .party_identifier import PartyIdentifier
from .typo_checker import TypoChecker

logger = logging.getLogger(__name__)

# 等级标记 → 枚举。LLM 的 reason 以「【高】…」开头，这是 prompt 约定
# （见 contract_reviewer 的企业合同风险管控框架）。
_SEVERITY_MARKERS: dict[str, str] = {
    "高": RevisionSeverity.HIGH,
    "中": RevisionSeverity.MEDIUM,
    "低": RevisionSeverity.LOW,
}

# 标题从 reason 正文里截到首个句末标点为止：左栏一行放不下整段理由，
# 展示「一句话结论」，完整依据仍在 reason 里供展开。
# 注意只停句末标点（。；）与追加说明的引导符（：（），不停顿号逗号——
# 「付款节点过早，钱先付完工期失控」是一整句结论，切在逗号上就丢了后半句。
_TITLE_BREAKS = "。；;：（(\n"
_TITLE_MAX = 200  # 与 ReviewRevision.title 的 max_length 对齐


def _parse_severity(reason: str) -> str:
    """从 '【高】付款节点过早…' 里取等级；取不到返回空串（前端按「未分级」展示）。"""
    text = (reason or "").lstrip()
    if not text.startswith("【"):
        return ""
    end = text.find("】")
    if end < 0:
        return ""
    return _SEVERITY_MARKERS.get(text[1:end].strip(), "")


def _make_title(reason: str) -> str:
    """剥掉等级标记与「依据：…」尾巴，截到首个断句符，得到一句话标题。"""
    text = (reason or "").lstrip()
    if text.startswith("【"):
        end = text.find("】")
        if end >= 0:
            text = text[end + 1 :]
    # 依据部分属于展开内容，不进标题
    for marker in ("依据：", "依据:"):
        idx = text.find(marker)
        if idx >= 0:
            text = text[:idx]
    for ch in _TITLE_BREAKS:
        idx = text.find(ch)
        if idx >= 0:
            text = text[:idx]
    return text.strip()[:_TITLE_MAX]


def _persist_revisions(task: ReviewTask, reviews: list[ReviewResult], *, kind: str = RevisionKind.RISK) -> int:
    """把审查产出落成结构化清单（复审台左栏的数据来源）。

    为什么必须有这一步：AI 修订原先只写进 docx 的 Track Changes，前端拿不到
    清单，也就无法「逐条采纳/驳回」。docx 里的痕迹是产物，这张表是决策载体。

    重跑语义（关键）：只清 ``pending`` 行，``accepted``/``rejected`` 一律保留。
    律师做出的判断不是 AI 的中间产物——AI 重跑一次就把它抹掉，等于逼律师
    每次重跑后从头复核一遍。

    ``kind`` 区分来源：AI 审查（risk）与从 docx 解析回流的人工修订（manual）。
    清理只针对**同来源**的 pending 行：重跑 AI 审查不能把律师的手工改动条目
    一并刷掉（反过来也一样）。

    待决策行不做「按原文去重」：AI 重跑可能对同一处给出修订后的建议（措辞
    变、条数变），旧的待决策建议本就是被 AI 自己推翻的，保留它们只会让律师
    在同一处看到互相矛盾的两条。已决策的才必顶保留。
    """
    if not reviews:
        return 0

    # 只删**同来源**的待决策旧行；已决策的、其他来源的一律保留
    ReviewRevision.objects.filter(task=task, status=RevisionStatus.PENDING, kind=kind).delete()

    # 序号接在既有行后面：人工条目不能从 0 开始，否则与 AI 条目序号撞车、
    # 排序时交错倒置
    existing = ReviewRevision.objects.filter(task=task)
    base = existing.aggregate(m=Max("ordinal"))["m"] if existing.exists() else None
    next_ordinal = (int(base) + 1) if base is not None else 0

    # 不按原文去重：AI 重跑本来就会重新给出一批建议（包括已被采纳的那处），
    # 全部进待决策，由律师看到最新的一批；已决策的旧行仍独立保留。
    rows: list[ReviewRevision] = []
    for offset, rev in enumerate(reviews):
        rows.append(
            ReviewRevision(
                task=task,
                ordinal=next_ordinal + offset,
                kind=kind,
                severity=_parse_severity(rev.reason),
                title=_make_title(rev.reason),
                original=rev.original or "",
                suggested=rev.suggested or "",
                reason=rev.reason or "",
                paragraph_index=getattr(rev, "paragraph_index", -1),
                status=RevisionStatus.PENDING,
            )
        )

    ReviewRevision.objects.bulk_create(rows)
    return len(rows)


def _ingest_manual_revisions(task: ReviewTask, docx_bytes: bytes) -> int:
    """把 docx 里的修订痕迹解析成人工条目落库。

    只收**非本系统作者**的痕迹：AI 自己写的修订（reviewer_name，默认
    「法穿SI」）已经在 `ReviewRevision` 里有一份（kind=risk），再当成人工
    条目收一次就会重得，而且律师会在左栏看到同一处两个条目。
    """
    from docx import Document as _Document

    from .docx_revision_reader import extract_revisions

    doc = _Document(BytesIO(docx_bytes))
    found = extract_revisions(doc)

    ai_author = (task.reviewer_name or "").strip()
    manual = [r for r in found if r.author.strip() != ai_author]
    if not manual:
        logger.info("人工修订回流: 无新条目（共 %d 条痕迹）", len(found))
        return 0

    # 转成 ReviewResult 同构：手动修改没有风险等级与法条依据。
    # 标题直接给改动内容（左栏一眼看出改了什么），而不是「律师手改」这种废话。
    converted = []
    for r in manual:
        if r.original and r.suggested:
            headline = f"改为「{r.suggested[:24]}」"
        elif r.original:
            headline = f"删除「{r.original[:24]}」"
        else:
            headline = f"新增「{r.suggested[:24]}」"
        converted.append(
            ReviewResult(
                original=r.original,
                suggested=r.suggested,
                # 【】留空等级：人工修订不分高中低，前端按「未分级」展示
                reason=f"【】{headline} · {r.author}",
                paragraph_index=r.paragraph_index,
            )
        )
    n = _persist_revisions(task, converted, kind=RevisionKind.MANUAL)
    logger.info("人工修订回流: %d 条落库（共 %d 条痕迹）", n, len(found))
    return n


def _resolve_task_file(file_path: str, error_message: str) -> Path:  # pragma: no cover
    """把落库路径解析为 MEDIA_ROOT 下的绝对路径（兼容存量绝对路径数据）。"""
    try:
        return to_media_abs(file_path)
    except ValidationException as e:
        raise ContractReviewError(error_message) from e


class ReviewService:  # pragma: no cover
    """合同审查主编排服务"""

    def __init__(self) -> None:  # pragma: no cover
        self._repository = ReviewTaskRepository()

    def upload_contract(  # pragma: no cover
        self,
        file: UploadedFile,
        user: object,
        model_name: str = "",
    ) -> ReviewTask:
        """上传合同：验证 → 保存 → 提取内容 → 识别甲乙方 → 创建任务"""
        filename = file.name or "unknown.docx"
        if not filename.lower().endswith(".docx"):
            raise ContractReviewError("仅支持 .docx 格式文件")

        # 保存文件（落库协议：media 相对路径）
        task_id = uuid.uuid4()
        safe_name = sanitize_upload_filename(filename)
        rel_path = f"{MediaEntity.CONTRACT_REVIEW_UPLOADS}/{task_id}_{safe_name}"
        saved_name = default_storage.save(rel_path, file)
        save_path = to_media_abs(saved_name)

        # 提取内容 + 识别甲乙方
        extractor = ContentExtractor()
        identifier = PartyIdentifier()
        try:
            paragraphs = extractor.extract_paragraphs(save_path)
            parties = identifier.identify_parties(paragraphs)
        except ExtractionError as e:
            task = self._repository.create(
                id=task_id,
                user=user,
                original_file=saved_name,
                status=TaskStatus.EXTRACTION_FAILED,
                error_message=str(e),
                model_name=model_name,
            )
            return task

        # 提取标题
        title_extractor = TitleExtractor()
        doc = Document(str(save_path))
        title = title_extractor.extract_title(doc)

        task = self._repository.create(
            id=task_id,
            user=user,
            original_file=saved_name,
            contract_title=title or filename.rsplit(".", 1)[0],
            party_a=parties.get("party_a", ""),
            party_b=parties.get("party_b", ""),
            party_c=parties.get("party_c", ""),
            party_d=parties.get("party_d", ""),
            status=TaskStatus.PARTIES_IDENTIFIED,
            model_name=model_name,
        )
        return task

    def confirm_party(  # pragma: no cover
        self,
        task_id: uuid.UUID,
        represented_party: str,
        user: object,
        reviewer_name: str = "",
        selected_steps: list[str] | None = None,
        party_overrides: dict[str, str] | None = None,
        client_note: str = "",
        client_demands: list[str] | None = None,
        our_client_name: str = "",
        stance_basis: str = "",
    ) -> ReviewTask:
        """确认代表方，提交异步审查任务"""
        task = self._repository.get_by_id_required(task_id)
        if task.status not in (
            TaskStatus.PARTIES_IDENTIFIED,
            TaskStatus.EXTRACTION_FAILED,
        ):
            raise ContractReviewError(f"当前状态 {task.status} 不允许确认代表方")

        # 默认全选
        default_steps = ["typo_check", "format_document", "contract_review", "review_report"]
        steps = selected_steps if selected_steps else default_steps

        reviewer_name = reviewer_name.strip() or "法穿SI"

        # 构建更新字段
        update_fields: dict[str, object] = {
            "represented_party": represented_party,
            "reviewer_name": reviewer_name,
            "status": TaskStatus.CONFIRMED,
            "selected_steps": steps,
            # 客户交代是审查方向的输入，不存下来等于律师白写
            "client_note": (client_note or "").strip(),
            "client_demands": client_demands or [],
            "our_client_name": (our_client_name or "").strip(),
            "stance_basis": (stance_basis or "").strip(),
        }
        # 允许用户手动修正识别错误的当事人名称
        if party_overrides:
            for key in ("party_a", "party_b", "party_c", "party_d"):
                if key in party_overrides:
                    update_fields[key] = party_overrides[key]

        self._repository.update(task_id, **update_fields)

        # 提交异步任务
        from apps.core.tasking import submit_task

        submit_task(
            "apps.contract_review.services.review.review_service.process_review",
            str(task_id),
            timeout=1800,
        )
        logger.info("已提交审查任务: %s", task_id)
        return self._repository.get_by_id(task_id)  # type: ignore[return-value]

    def get_task_status(self, task_id: uuid.UUID) -> ReviewTask:  # pragma: no cover
        """查询任务状态"""
        return self._repository.get_by_id_required(task_id)

    def get_result_file(self, task_id: uuid.UUID) -> Path:  # pragma: no cover
        """获取结果文件路径"""
        task = self._repository.get_by_id_required(task_id)
        if task.status != TaskStatus.COMPLETED:
            raise ContractReviewError("任务尚未完成")
        path = _resolve_task_file(task.output_file, "结果文件不存在")
        if not path.exists():
            raise ContractReviewError("结果文件不存在")
        return path

    def get_original_file(self, task_id: uuid.UUID) -> Path:  # pragma: no cover
        """获取原始上传文件路径"""
        task = self._repository.get_by_id_required(task_id)
        path = _resolve_task_file(task.original_file, "原始文件不存在")
        if not path.exists():
            raise ContractReviewError("原始文件不存在")
        return path

    def adopt_edited_output(self, task_id: uuid.UUID, data: bytes, *, version_tag: int) -> str:
        """收养 OnlyOffice 保存回写的编辑产物（律师画布内编辑落盘）。

        与审查流水线的产物落盘同一模式：default_storage + CONTRACT_REVIEW_OUTPUT；
        旧产物物理文件保留（调用方在 EditorSession.prev_output_file 留档路径），
        出问题可手工找回。返回新的落库相对路径。
        """
        task = self._repository.get_by_id_required(task_id)
        if task.status != TaskStatus.COMPLETED:
            raise ContractReviewError("任务尚未完成，不能收养编辑产物")
        old_name = Path(task.output_file or "").stem or "review"
        output_name = f"{old_name}-律编v{version_tag}.docx"
        output_rel = f"{MediaEntity.CONTRACT_REVIEW_OUTPUT}/{sanitize_upload_filename(output_name)}"
        saved_name = default_storage.save(output_rel, ContentFile(data))
        self._repository.update(task_id, output_file=saved_name)
        logger.info("OnlyOffice 编辑产物已收养: task=%s -> %s", task_id, saved_name)

        # 人工修订回流清单：律师在编辑器里自己改的内容存成 w:ins/w:del，
        # 左栏原本看不到，他改完一份合同清单里还是「待决策 12 条」。
        # 回流失败不能阻断保存——文件已经落盘，丢了清单比丢了正文轻。
        try:
            _ingest_manual_revisions(task, data)
        except Exception:
            logger.warning("人工修订回流失败（不影响保存）: task=%s", task_id, exc_info=True)

        # 版本链：每次保存是一个新版本节点（律师编辑版）
        try:
            from apps.contract_review.models.document_version import VersionKind

            from .version_chain import record_version

            record_version(
                task_id=task_id,
                kind=VersionKind.LAWYER,
                file_path=saved_name,
                file_name=output_name,
                author=task.reviewer_name or "律师",
                note=f"编辑器保存 v{version_tag}",
            )
        except Exception:
            logger.warning("版本链记录失败（不影响保存）: task=%s", task_id, exc_info=True)

        return saved_name


def process_review(task_id_str: str) -> None:  # pragma: no cover
    """异步执行审查流水线（由 Django-Q2 调用）"""
    task_id = uuid.UUID(task_id_str)
    repository = ReviewTaskRepository()
    task = repository.get_by_id(task_id)
    if task is None:
        logger.warning("任务不存在，跳过: %s", task_id)
        return
    # 防止重复执行：已完成/已失败的任务直接跳过
    if task.status in (TaskStatus.COMPLETED, TaskStatus.FAILED):
        logger.warning("任务已终态 (%s)，跳过: %s", task.status, task_id)
        return
    repository.update(task_id, status=TaskStatus.PROCESSING)

    try:
        original_path = to_media_abs(task.original_file)
        doc = Document(str(original_path))
        extractor = ContentExtractor()
        extraction = extractor.extract_with_mapping(original_path)
        paragraphs = extraction.paragraphs

        llm: LLMService = get_llm_service()
        revision_tool = DocxRevisionTool()

        # 启用修订模式（Track Changes）
        enable_track_changes(doc)

        steps = getattr(task, "selected_steps", None) or [
            "typo_check",
            "format_document",
            "contract_review",
            "review_report",
        ]

        # Step 1: 标题提取（始终执行）
        _update_step(repository, task, ProcessStep.TITLE_EXTRACTION)
        title_extractor = TitleExtractor()
        title = title_extractor.extract_title(doc)
        if title:
            repository.update(task.id, contract_title=title)

        # Step 2: 错别字检测
        if "typo_check" in steps:
            _update_step(repository, task, ProcessStep.TYPO_CHECK)
            typo_checker = TypoChecker(llm)
            typos = typo_checker.check_typos(paragraphs, model_name=task.model_name)
            typo_applied = 0
            for typo in typos:
                applied = _apply_to_any_paragraph(
                    doc, revision_tool, typo.original, typo.corrected, author=task.reviewer_name
                )
                if applied:
                    typo_applied += 1
            logger.info("错别字修订: %d/%d 处成功", typo_applied, len(typos))

        # Step 3: 合同审查 + 评估报告（并行执行，两者互不依赖）
        need_review = "contract_review" in steps
        need_report = "review_report" in steps
        reviewer = ContractReviewer(llm)

        if need_review and need_report:
            _update_step(repository, task, ProcessStep.CONTRACT_REVIEW)
            reviews: list[ReviewResult] = []
            report = ""
            # 并行子任务失败不能只打日志就吞掉：全失败要标 FAILED，部分失败要留痕
            failed_kinds: list[str] = []
            failure_detail: list[str] = []

            def _run_review() -> list[ReviewResult]:  # pragma: no cover
                return reviewer.review_contract(
                    paragraphs,
                    task.represented_party,
                    task.party_a,
                    task.party_b,
                    model_name=task.model_name,
                    client_note=task.client_note,
                    client_demands=task.client_demands,
                )

            def _run_report() -> str:  # pragma: no cover
                return reviewer.generate_report(
                    paragraphs,
                    task.represented_party,
                    task.party_a,
                    task.party_b,
                    model_name=task.model_name,
                )

            with ThreadPoolExecutor(max_workers=2) as pool:
                futures: dict[Future[object], str] = {
                    pool.submit(_run_review): "review",
                    pool.submit(_run_report): "report",
                }
                for future in as_completed(futures):
                    kind = futures[future]
                    try:
                        if kind == "review":
                            reviews = future.result()  # type: ignore[assignment]
                        else:
                            report = future.result()
                    except Exception as e:
                        failed_kinds.append(kind)
                        failure_detail.append(f"{kind}: {e}")
                        logger.exception("并行任务 %s 失败", kind)

            # reviews 与 report 全为空说明并行任务毫无产出：抛错走外层 FAILED 分支，
            # 禁止继续往下产出空文档并虚报 COMPLETED
            if failed_kinds and not reviews and not report:
                raise ContractReviewError("合同审查与评估报告并行任务均失败: " + "; ".join(failure_detail))

            # 部分失败：任务仍可完成，但把失败明细写进 error_message 留痕（不覆盖成功产物）
            if failed_kinds:
                repository.update(
                    task.id,
                    error_message=f"{len(failed_kinds)} 项并行子任务失败: " + "; ".join(failure_detail),
                )

            review_applied = 0
            for rev in reviews:
                applied = _apply_to_any_paragraph(
                    doc, revision_tool, rev.original, rev.suggested, author=task.reviewer_name
                )
                if applied:
                    review_applied += 1
            logger.info("合同审查修订: %d/%d 处成功", review_applied, len(reviews))

            # 落结构化清单：docx 里的是痕迹（产物），这张表是律师决策的载体
            persisted = _persist_revisions(task, reviews)
            logger.info("审查清单落库: %d 条待决策", persisted)

            if report:
                repository.update(task.id, review_report=report)
                logger.info("评估报告已生成 (%d 字)", len(str(report)))

        elif need_review:
            _update_step(repository, task, ProcessStep.CONTRACT_REVIEW)
            reviews_only = reviewer.review_contract(
                paragraphs,
                task.represented_party,
                task.party_a,
                task.party_b,
                model_name=task.model_name,
                client_note=task.client_note,
                client_demands=task.client_demands,
            )
            review_applied = 0
            for rev in reviews_only:
                applied = _apply_to_any_paragraph(
                    doc, revision_tool, rev.original, rev.suggested, author=task.reviewer_name
                )
                if applied:
                    review_applied += 1
            logger.info("合同审查修订: %d/%d 处成功", review_applied, len(reviews_only))

            persisted = _persist_revisions(task, reviews_only)
            logger.info("审查清单落库: %d 条待决策", persisted)

        elif need_report:
            report_only = reviewer.generate_report(
                paragraphs,
                task.represented_party,
                task.party_a,
                task.party_b,
                model_name=task.model_name,
            )
            if report_only:
                repository.update(task.id, review_report=report_only)
                logger.info("评估报告已生成 (%d 字)", len(report_only))

        # Step 4: 格式标准化
        if "format_document" in steps:
            _update_step(repository, task, ProcessStep.FORMAT_DOCUMENT)
            DocxFormatter().format_document(doc)

            _update_step(repository, task, ProcessStep.PAGE_NUMBERING)
            PageNumbering().standardize(doc)

            _update_step(repository, task, ProcessStep.HEADING_NUMBERING)
            HeadingNumbering(llm).apply_numbering(doc, model_name=task.model_name)

        # 保存输出文件（经 default_storage 落盘，落库存相对路径）
        output_name = generate_output_filename(
            task.contract_title,
            task_id=str(task.id),
        )
        output_rel = f"{MediaEntity.CONTRACT_REVIEW_OUTPUT}/{sanitize_upload_filename(output_name)}"
        buffer = BytesIO()
        doc.save(buffer)
        saved_name = default_storage.save(output_rel, ContentFile(buffer.getvalue()))

        # 版本链：链头（客户来稿）+ AI 产物节点。链头先确保存在，
        # 否则第一个 AI 版会成为链头，看不出「从哪份来稿改出来的」。
        try:
            from apps.contract_review.models.document_version import VersionKind

            from .version_chain import ensure_original_version, record_version

            ensure_original_version(
                task_id=task.id,
                file_path=task.original_file,
                file_name=Path(task.original_file).name,
            )
            record_version(
                task_id=task.id,
                kind=VersionKind.AI,
                file_path=saved_name,
                file_name=output_name,
                author=task.reviewer_name or "法穿SI",
                note="AI 审查产物",
            )
        except Exception:
            logger.warning("版本链记录失败（不影响完成状态）: task=%s", task_id, exc_info=True)

        repository.update(
            task.id,
            output_file=saved_name,
            status=TaskStatus.COMPLETED,
            current_step="",
        )
        logger.info("审查完成: %s -> %s", task_id, saved_name)

    except Exception as e:
        logger.exception("审查任务失败: %s", task_id)
        repository.update(
            task_id,
            status=TaskStatus.FAILED,
            error_message=str(e),
        )


def _update_step(repository: ReviewTaskRepository, task: ReviewTask, step: str) -> None:  # pragma: no cover
    repository.update(task.id, current_step=step)


def _apply_to_any_paragraph(  # pragma: no cover
    doc: DocumentType, tool: DocxRevisionTool, original: str, replacement: str, author: str = ""
) -> bool:
    """遍历所有段落查找原文并应用修订，不依赖 LLM 返回的段落索引"""
    for para in doc.paragraphs:
        if original in para.text:
            if tool.apply_revision(para, original, replacement, author=author or None):
                return True
    return False
