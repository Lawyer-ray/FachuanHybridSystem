"""安全审计（2026-10-19）回归测试。

覆盖：
- P1：docxtpl 占位符值 OOXML 注入（autoescape 收口，document.xml 断言 + python-docx 打开）
- P2：ZIP 条目名清洗（safe_arcname / ZipPackager / FolderGenerationService / 批量填充打包）
- P2：post_delete 物理文件清理收敛在 MEDIA_ROOT 内
- P2：私有 docx 模板根目录白名单
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest

from apps.documents.services.generation.path_utils import safe_arcname

# 审计给定的注入载荷：尝试逃逸 <w:t> 文本节点并注入域代码（DDE）
INJECTION_PAYLOAD = '</w:t></w:r><w:fldSimple w:instr="DDE" dirty="true"/>'


def _make_minimal_docx(path: Path, body: str) -> None:
    """构造仅含 document.xml 的最小 docx，用于验证 docxtpl 渲染安全行为。"""
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body>{body}</w:body></w:document>",
        )


def _document_xml_of(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


class TestDocxRenderXmlInjection:
    """P1：占位符值注入 docx XML（安全审计 2026-10-19）。"""

    def test_pipeline_renderer_escapes_ooxml_injection_payload(self, tmp_path: Path) -> None:
        """生成管线 DocxRenderer 渲染后，载荷被 XML 转义且文档可被 python-docx 正常打开。"""
        from apps.documents.services.generation.pipeline.renderer import DocxRenderer

        tpl = tmp_path / "tpl.docx"
        _make_minimal_docx(tpl, "<w:p><w:r><w:t>客户：{{ 客户名称 }}</w:t></w:r></w:p>")

        rendered = DocxRenderer().render(str(tpl), {"客户名称": INJECTION_PAYLOAD})

        xml = _document_xml_of(rendered)
        # 注入载荷不得以原样出现在 document.xml 中
        assert INJECTION_PAYLOAD not in xml
        assert '<w:fldSimple w:instr="DDE"' not in xml
        # 转义后的实体必须出现
        assert "&lt;/w:t&gt;" in xml
        assert "&lt;w:fldSimple" in xml
        # 文档结构完好，可被 python-docx 正常打开
        from docx import Document

        doc = Document(io.BytesIO(rendered))
        # 文本语义不降级：段落文本保留原始字面量
        assert INJECTION_PAYLOAD in doc.paragraphs[0].text

    def test_sandbox_env_render_escapes_payload_for_direct_render_sites(self, tmp_path: Path) -> None:
        """非管线的直连渲染点（preservation 等）统一经 SANDBOXED_JINJA_ENV，同样被转义覆盖。"""
        from docxtpl import DocxTemplate

        from apps.documents.services.placeholders.fallback import SANDBOXED_JINJA_ENV, build_docx_render_context

        tpl = tmp_path / "tpl.docx"
        _make_minimal_docx(
            tpl,
            "<w:p><w:r><w:t>案由：{{ 案由 }}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>{%p for item in 卷内目录 %}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>{{ item.材料名称 }}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>{%p endfor %}</w:t></w:r></w:p>",
        )

        doc = DocxTemplate(str(tpl))
        context = build_docx_render_context(
            doc=doc,
            context={"案由": INJECTION_PAYLOAD, "卷内目录": [{"序号": 1, "材料名称": INJECTION_PAYLOAD, "页码": "1"}]},
        )
        doc.render(context, jinja_env=SANDBOXED_JINJA_ENV)
        buffer = io.BytesIO()
        doc.save(buffer)

        xml = _document_xml_of(buffer.getvalue())
        assert INJECTION_PAYLOAD not in xml
        assert "&lt;/w:t&gt;" in xml

    def test_listing_hard_break_and_no_double_escape(self, tmp_path: Path) -> None:
        """docxtpl Listing（含 \\n 硬换行）与普通字符串混用时：Listing 不被二次转义，换行仍转 <w:br/>。"""
        from docxtpl import DocxTemplate, Listing

        from apps.documents.services.placeholders.fallback import SANDBOXED_JINJA_ENV, build_docx_render_context

        tpl = tmp_path / "tpl.docx"
        _make_minimal_docx(
            tpl,
            "<w:p><w:r><w:t>日志：{{ 律师工作日志内容 }}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>名称：{{ 合同名称 }}</w:t></w:r></w:p>",
        )

        doc = DocxTemplate(str(tpl))
        context = build_docx_render_context(
            doc=doc,
            context={
                "律师工作日志内容": Listing("第一行" + INJECTION_PAYLOAD + "\n第二行"),
                "合同名称": "A&B 公司",
            },
        )
        doc.render(context, jinja_env=SANDBOXED_JINJA_ENV)
        buffer = io.BytesIO()
        doc.save(buffer)

        xml = _document_xml_of(buffer.getvalue())
        # Listing 原样放行：换行转为硬换行、不出现二次转义
        assert "<w:br/>" in xml
        assert "&amp;lt;" not in xml
        # 普通字符串中的 & 被正确转义为 XML 实体
        assert "A&amp;B 公司" in xml


class TestSafeArcname:
    """P2：ZIP 条目名清洗（拒 .. 与绝对路径形态）。"""

    def test_drops_dot_dot_segments(self) -> None:
        assert safe_arcname("../../etc/passwd") == "etc/passwd"
        assert safe_arcname("a/../b") == "a/b"
        assert safe_arcname("a/./b") == "a/b"

    def test_strips_absolute_and_drive_prefix(self) -> None:
        assert safe_arcname("/etc/passwd") == "etc/passwd"
        assert safe_arcname("C:\\evil\\x.docx") == "evil/x.docx"
        assert safe_arcname("c:/evil/x.docx") == "evil/x.docx"

    def test_backwards_compatible_benign_names(self) -> None:
        assert safe_arcname("contracts/file.pdf") == "contracts/file.pdf"
        assert safe_arcname(r"contracts\file.pdf") == "contracts/file.pdf"
        assert safe_arcname("a//b") == "a/b"
        assert safe_arcname("/a/b") == "a/b"
        assert safe_arcname("") == ""

    def test_only_traversal_name_falls_back_to_empty(self) -> None:
        assert safe_arcname("..") == ""
        assert safe_arcname("./") == ""


class TestZipEntriesSanitized:
    """P2：骨架节点名与文件路径进 ZIP 前统一清洗。"""

    def test_packager_sanitizes_folder_skeleton_and_document_paths(self) -> None:
        from apps.documents.services.generation.pipeline.packager import ZipPackager

        structure = {
            "name": "../root",
            "children": [
                {"name": "..", "children": []},
                {"name": "/etc/passwd", "children": [{"name": "正常子目录", "children": []}]},
            ],
        }
        documents = [
            ("../../etc", b"content1", "恶意../../文件.docx"),
            ("子目录", b"content2", "C:\\绝对路径.txt"),
        ]

        zip_bytes = ZipPackager().create(structure, documents)
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            names = zf.namelist()

        for name in names:
            assert not name.startswith("/")
            assert ".." not in Path(name).parts
            assert "C:" not in name
        # 正常目录结构保留
        assert any("正常子目录" in name for name in names)
        # 文件条目路径被清洗（.. 段剔除、盘符剥除），保留可见文件名
        assert "root/etc/恶意../文件.docx" in names
        assert "root/子目录/绝对路径.txt" in names

    def test_folder_generation_service_skeleton_sanitized(self) -> None:
        from apps.documents.services.generation.folder_generation_service import FolderGenerationService

        service = FolderGenerationService(contract_service=MagicMock(), folder_binding_service=MagicMock())
        structure = {
            "name": "案件卷",
            "children": [
                {"name": "../escape", "children": []},
                {"name": "/绝对路径子目录", "children": []},
                {"name": "正常子目录", "children": []},
            ],
        }
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as zf:
            service._create_folders_in_zip(zf, structure, "")
        with zipfile.ZipFile(buffer) as zf:
            names = zf.namelist()

        for name in names:
            assert ".." not in Path(name).parts
            assert not name.startswith("/")
        assert "案件卷/escape/" in names
        assert "案件卷/绝对路径子目录/" in names
        assert "案件卷/正常子目录/" in names

    def test_filling_service_pack_to_zip_sanitizes_arcname(self, tmp_path: Path) -> None:
        from apps.documents.services.external_template.filling_service import FillingService

        record_rel = "documents/external_filled/1/fill_1.docx"
        file_abs = tmp_path / record_rel
        file_abs.parent.mkdir(parents=True)
        file_abs.write_bytes(b"filled")
        records = [
            SimpleNamespace(
                id=1,
                case_id=1,
                batch_task_id=9,
                file_path=record_rel,
                original_output_name="../../../../tmp/pwn.docx",
            )
        ]

        with patch("apps.documents.services.external_template.filling_service.settings") as mock_settings:
            mock_settings.MEDIA_ROOT = str(tmp_path)
            service = FillingService(placeholder_registry=MagicMock())
            zip_relative = service._pack_to_zip(records)  # type: ignore[arg-type]

        zip_abs = tmp_path / zip_relative
        with zipfile.ZipFile(zip_abs) as zf:
            names = zf.namelist()

        assert names == ["tmp/pwn.docx"]
        for name in names:
            assert ".." not in Path(name).parts


class TestPostDeleteFileCleanupConfinement:
    """P2：post_delete unlink 收敛在 MEDIA_ROOT 内。"""

    def _media_root(self, tmp_path: Path) -> Path:
        media = tmp_path / "media"
        media.mkdir()
        return media

    def test_deletes_file_inside_media_root(self, tmp_path: Path) -> None:
        from apps.documents.signals import _delete_charfield_file

        media = self._media_root(tmp_path)
        target = media / "documents" / "a.docx"
        target.parent.mkdir(parents=True)
        target.write_bytes(b"x")

        with patch("apps.documents.signals.settings") as mock_settings:
            mock_settings.MEDIA_ROOT = str(media)
            _delete_charfield_file("documents/a.docx")

        assert not target.exists()

    def test_skips_traversal_outside_media_root(self, tmp_path: Path) -> None:
        from apps.documents.signals import _delete_charfield_file

        media = self._media_root(tmp_path)
        outside = tmp_path / "outside.docx"
        outside.write_bytes(b"secret")

        with patch("apps.documents.signals.settings") as mock_settings:
            mock_settings.MEDIA_ROOT = str(media)
            _delete_charfield_file("../outside.docx")

        assert outside.exists()

    def test_skips_absolute_path_outside_media_root(self, tmp_path: Path) -> None:
        from apps.documents.signals import _delete_charfield_file

        media = self._media_root(tmp_path)
        outside = tmp_path / "elsewhere" / "b.docx"
        outside.parent.mkdir(parents=True)
        outside.write_bytes(b"secret")

        with patch("apps.documents.signals.settings") as mock_settings:
            mock_settings.MEDIA_ROOT = str(media)
            _delete_charfield_file(str(outside))

        assert outside.exists()


class TestPrivateDocxRootWhitelist:
    """P2：私有 docx 模板根目录白名单（MEDIA_ROOT / backend 项目目录之内）。"""

    def test_root_inside_backend_dir_accepted(self, tmp_path: Path) -> None:
        from apps.documents.admin.template_admin_views_mixin import _normalize_private_docx_root

        backend_root = tmp_path / "backend"
        candidate = backend_root / "private_templates"
        candidate.mkdir(parents=True)

        with (
            patch("django.conf.settings") as mock_settings,
            patch("apps.documents.storage.settings") as mock_storage_settings,
        ):
            mock_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_settings.MEDIA_ROOT = str(tmp_path / "media")
            mock_storage_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_storage_settings.MEDIA_ROOT = str(tmp_path / "media")
            result = _normalize_private_docx_root(str(candidate))
            assert Path(result).resolve() == candidate.resolve()

    def test_root_inside_media_root_accepted(self, tmp_path: Path) -> None:
        from apps.documents.admin.template_admin_views_mixin import _normalize_private_docx_root

        backend_root = tmp_path / "backend"
        media_root = tmp_path / "media"
        candidate = media_root / "docx_private"
        candidate.mkdir(parents=True)

        with (
            patch("django.conf.settings") as mock_settings,
            patch("apps.documents.storage.settings") as mock_storage_settings,
        ):
            mock_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_settings.MEDIA_ROOT = str(media_root)
            mock_storage_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_storage_settings.MEDIA_ROOT = str(media_root)
            result = _normalize_private_docx_root(str(candidate))
            assert Path(result).resolve() == candidate.resolve()

    def test_root_outside_whitelist_rejected(self, tmp_path: Path) -> None:
        from apps.documents.admin.template_admin_views_mixin import _normalize_private_docx_root

        backend_root = tmp_path / "backend"
        (backend_root / "apiSystem").mkdir(parents=True)

        with (
            patch("django.conf.settings") as mock_settings,
            patch("apps.documents.storage.settings") as mock_storage_settings,
        ):
            mock_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_settings.MEDIA_ROOT = str(tmp_path / "media")
            mock_storage_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_storage_settings.MEDIA_ROOT = str(tmp_path / "media")
            # tmp_path 本身存在且是目录，但不在白名单内 → 拒绝
            with pytest.raises(ValueError, match="越界"):
                _normalize_private_docx_root(str(tmp_path))

    @pytest.mark.parametrize("evil_root", ["/", "/etc", "/var"])
    def test_system_top_level_dirs_rejected(self, evil_root: str) -> None:
        from apps.documents.admin.template_admin_views_mixin import _normalize_private_docx_root

        if not Path(evil_root).is_dir():
            pytest.skip(f"{evil_root} 在当前平台不存在")
        # 使用真实 settings：/、/etc 等顶级目录必然在白名单之外
        with pytest.raises(ValueError, match="越界"):
            _normalize_private_docx_root(evil_root)

    def test_storage_whitelist_helper(self, tmp_path: Path) -> None:
        from apps.documents.storage import is_allowed_private_docx_root

        backend_root = tmp_path / "backend"
        with patch("apps.documents.storage.settings") as mock_settings:
            mock_settings.BASE_DIR = str(backend_root / "apiSystem")
            mock_settings.MEDIA_ROOT = str(tmp_path / "media")
            assert is_allowed_private_docx_root(backend_root / "sub" / "root") is True
            assert is_allowed_private_docx_root(tmp_path / "media" / "x") is True
            assert is_allowed_private_docx_root(Path("/etc")) is False
