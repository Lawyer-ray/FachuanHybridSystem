"""placeholders 顶层服务补覆盖：placeholder_usage_service / placeholder_service / placeholder_admin_service。

- usage 服务：模板签名缓存键 + python-docx 兜底抽取（真实 docx 用 tmp_path 构造）
- 占位符服务：注册校验与按 ID/键查询、软删除（真实 ORM）
- admin 服务：查询集过滤逻辑（MagicMock 查询集）
"""

from __future__ import annotations

from datetime import datetime
from unittest.mock import MagicMock, patch

import pytest

USAGE_SVC = "apps.documents.services.placeholders.placeholder_usage_service"


# ── PlaceholderUsageService ─────────────────────────────────────────


class TestPlaceholderUsageServiceSignature:
    def _svc(self):
        from apps.documents.services.placeholders.placeholder_usage_service import PlaceholderUsageService

        return PlaceholderUsageService()

    def _patch_aggregate(self, max_updated_at, count):
        manager = MagicMock()
        manager.filter.return_value.aggregate.return_value = {"max_updated_at": max_updated_at, "count": count}
        return patch(f"{USAGE_SVC}.DocumentTemplate", objects=manager)

    def test_signature_with_datetime(self):
        dt = datetime(2026, 1, 2, 3, 4, 5)
        with self._patch_aggregate(dt, 7):
            assert self._svc()._get_template_signature() == (int(dt.timestamp()), 7, "v1")

    def test_signature_without_datetime(self):
        with self._patch_aggregate(None, 0):
            assert self._svc()._get_template_signature() == (0, 0, "v1")

    def test_cache_key_contains_signature(self):
        dt = datetime(2026, 1, 2, 3, 4, 5)
        with self._patch_aggregate(dt, 7):
            key = self._svc()._get_cache_key()
        assert key.startswith("documents:placeholder_usage:")
        assert key.endswith(":7:v1")


class TestPlaceholderUsageServiceExtraction:
    def _svc(self):
        from apps.documents.services.placeholders.placeholder_usage_service import PlaceholderUsageService

        return PlaceholderUsageService()

    def _make_docx(self, path) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("合同编号：{{ 合同编号 }}，甲方：{{甲方名称}}")
        doc.add_paragraph("无占位符段落")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "{{ 乙方名称 }}"
        table.rows[0].cells[1].text = "{{合同编号}}"  # 与段落重复
        doc.save(str(path))

    def test_extract_by_python_docx(self, tmp_path):
        docx_file = tmp_path / "tpl.docx"
        self._make_docx(docx_file)
        keys = self._svc()._extract_by_python_docx(str(docx_file))
        assert keys == {"合同编号", "甲方名称", "乙方名称"}

    def test_extract_by_python_docx_missing_file(self, tmp_path):
        assert self._svc()._extract_by_python_docx(str(tmp_path / "missing.docx")) == set()

    def test_iter_doc_texts_paragraphs_and_tables(self):
        para = MagicMock()
        para.text = "段落A"
        cell = MagicMock()
        cell.text = "单元格B"
        row = MagicMock()
        row.cells = [cell]
        table = MagicMock()
        table.rows = [row]
        doc = MagicMock()
        doc.paragraphs = [para]
        doc.tables = [table]

        texts = self._svc()._iter_doc_texts(doc)
        assert texts == ["段落A", "单元格B"]

    def test_iter_doc_texts_broken_doc_returns_partial(self):
        class BrokenParas:
            @property
            def paragraphs(self):
                raise RuntimeError("broken")

        doc = BrokenParas()
        doc.tables = []  # type: ignore[attr-defined]
        assert self._svc()._iter_doc_texts(doc) == []

    def test_placeholder_pattern(self):
        from apps.documents.services.placeholders.placeholder_usage_service import PlaceholderUsageService

        pattern = PlaceholderUsageService._PLACEHOLDER_PATTERN
        assert pattern.findall("{{ 键A }}和{{键B}}") == ["键A", "键B"]
        assert pattern.findall("无占位符") == []


# ── PlaceholderService ──────────────────────────────────────────────


class TestPlaceholderServiceRegister:
    def _svc(self):
        from apps.documents.services.placeholders.placeholder_service import PlaceholderService

        return PlaceholderService()

    def test_register_blank_key_raises(self):
        from apps.core.exceptions import ValidationException

        with pytest.raises(ValidationException):
            self._svc().register_placeholder("", "显示名")

    def test_register_blank_display_name_raises(self):
        from apps.core.exceptions import ValidationException

        with pytest.raises(ValidationException):
            self._svc().register_placeholder("键", "")


@pytest.mark.django_db
class TestPlaceholderServiceQueries:
    def _svc(self):
        from apps.documents.services.placeholders.placeholder_service import PlaceholderService

        return PlaceholderService()

    def _create(self, key="键A") -> object:
        from apps.documents.models import Placeholder

        return Placeholder.objects.create(key=key, display_name=f"显示-{key}")

    def test_get_placeholder_by_id(self):
        placeholder = self._create()
        assert self._svc().get_placeholder_by_id(placeholder.id).key == "键A"  # type: ignore[attr-defined]

    def test_get_placeholder_by_id_missing_raises(self):
        from apps.core.exceptions import NotFoundError

        with pytest.raises(NotFoundError):
            self._svc().get_placeholder_by_id(99999)

    def test_get_placeholder_by_key(self):
        self._create("键B")
        assert self._svc().get_placeholder_by_key("键B").display_name == "显示-键B"

    def test_get_placeholder_by_key_missing_raises(self):
        from apps.core.exceptions import NotFoundError

        with pytest.raises(NotFoundError):
            self._svc().get_placeholder_by_key("不存在的键")

    def test_delete_placeholder_soft_deactivates(self):
        placeholder = self._create("键C")
        assert self._svc().delete_placeholder(placeholder.id) is True  # type: ignore[attr-defined]
        placeholder.refresh_from_db()  # type: ignore[attr-defined]
        assert placeholder.is_active is False  # type: ignore[attr-defined]

    def test_delete_placeholder_missing_raises(self):
        from apps.core.exceptions import NotFoundError

        with pytest.raises(NotFoundError):
            self._svc().delete_placeholder(99999)


# ── PlaceholderAdminService ─────────────────────────────────────────


class TestPlaceholderAdminService:
    def _svc(self):
        from apps.documents.services.placeholders.placeholder_admin_service import PlaceholderAdminService

        return PlaceholderAdminService()

    def test_get_filtered_queryset(self):
        base_qs = MagicMock()
        result = self._svc().get_filtered_queryset(base_qs, ["k1", "k2"])
        base_qs.filter.assert_called_once_with(key__in=["k1", "k2"])
        assert result is base_qs.filter.return_value

    def _usage_map(self):
        return {
            "仅合同": {"contract"},
            "仅案件": {"case"},
            "双用": {"contract", "case"},
            "其他": {"archive"},
        }

    def test_filter_by_usage_contract(self):
        qs = MagicMock()
        result = self._svc().filter_by_usage(qs, "contract", self._usage_map())
        qs.filter.assert_called_once_with(key__in={"仅合同"})
        assert result is qs.filter.return_value

    def test_filter_by_usage_case(self):
        qs = MagicMock()
        self._svc().filter_by_usage(qs, "case", self._usage_map())
        qs.filter.assert_called_once_with(key__in={"仅案件"})

    def test_filter_by_usage_both(self):
        qs = MagicMock()
        self._svc().filter_by_usage(qs, "both", self._usage_map())
        qs.filter.assert_called_once_with(key__in={"双用"})

    def test_filter_by_usage_unused_excludes_used(self):
        qs = MagicMock()
        self._svc().filter_by_usage(qs, "unused", self._usage_map())
        qs.exclude.assert_called_once_with(key__in=set(self._usage_map().keys()))

    def test_filter_by_usage_unknown_returns_queryset(self):
        qs = MagicMock()
        result = self._svc().filter_by_usage(qs, "anything", self._usage_map())
        qs.filter.assert_not_called()
        qs.exclude.assert_not_called()
        assert result is qs
