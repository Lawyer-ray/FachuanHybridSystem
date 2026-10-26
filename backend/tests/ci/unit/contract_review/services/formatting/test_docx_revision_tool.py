"""Tests for contract_review.services.formatting.docx_revision_tool."""

from __future__ import annotations

from unittest.mock import MagicMock

import pytest

from apps.contract_review.services.formatting.docx_revision_tool import (
    DocxRevisionTool,
    _create_del,
    _create_ins,
    _make_run,
    _next_rev_id,
    enable_track_changes,
)


class TestNextRevId:
    def test_returns_string(self):
        result = _next_rev_id()
        assert isinstance(result, str)

    def test_increments(self):
        id1 = _next_rev_id()
        id2 = _next_rev_id()
        assert int(id2) > int(id1)


class TestMakeRun:
    def test_creates_run_element(self):
        from lxml import etree

        source = MagicMock()
        source._element = MagicMock()
        source._element.find.return_value = None
        result = _make_run("text", source)
        assert result is not None

    def test_with_del_text_tag(self):
        source = MagicMock()
        source._element = MagicMock()
        source._element.find.return_value = None
        result = _make_run("deleted", source, tag="w:delText")
        assert result is not None


class TestCreateDel:
    def test_creates_del_element(self):
        from lxml import etree

        source = MagicMock()
        source._element = MagicMock()
        source._element.find.return_value = None
        result = _create_del("old text", "author", "2024-01-01T00:00:00Z", source)
        assert result is not None


class TestCreateIns:
    def test_creates_ins_element(self):
        source = MagicMock()
        source._element = MagicMock()
        source._element.find.return_value = None
        result = _create_ins("new text", "author", "2024-01-01T00:00:00Z", source)
        assert result is not None


class TestDocxRevisionTool:
    def test_enable_track_changes(self):
        doc = MagicMock()
        settings_elem = MagicMock()
        doc.settings.element = settings_elem
        settings_elem.findall.return_value = []
        enable_track_changes(doc)
        settings_elem.append.assert_called_once()

    def test_apply_revision_not_found(self):
        tool = DocxRevisionTool()
        para = MagicMock()
        para.text = "no match here"
        para.runs = [MagicMock(text="no match here")]
        result = tool.apply_revision(para, "original", "replacement")
        assert result is False

    def test_apply_revision_single_run_match(self):
        """Test that single run match finds the text and returns True."""
        tool = DocxRevisionTool()
        # Create a real docx paragraph to test with
        from docx import Document

        doc = Document()
        para = doc.add_paragraph("prefix original suffix")
        result = tool.apply_revision(para, "original", "replacement")
        assert result is True
        # Verify del and ins elements were inserted
        xml_str = para._element.xml
        assert "w:del" in xml_str
        assert "w:ins" in xml_str

    def test_apply_revision_empty_runs(self):
        tool = DocxRevisionTool()
        para = MagicMock()
        para.text = ""
        para.runs = []
        result = tool.apply_revision(para, "original", "replacement")
        assert result is False

    def test_apply_revision_cross_run(self):
        """Test cross-run matching with real docx elements."""
        tool = DocxRevisionTool()
        from docx import Document

        doc = Document()
        para = doc.add_paragraph()
        # Add text split across runs
        run1 = para.add_run("hello ")
        run2 = para.add_run("world")
        result = tool.apply_revision(para, "hello world", "new text")
        assert result is True
        xml_str = para._element.xml
        assert "w:del" in xml_str
        assert "w:ins" in xml_str

    def test_apply_revision_cross_run_inserted_before_trailing_run(self):
        """跨 run 修订块必须插在被替换首 run 的原位，而非段落末尾。

        构造 4-run 段落（前缀/中间/尾巴/末尾），对「中间+尾巴」跨 run 文本做修订：
        del/ins 块应位于「前缀」run 之后、「末尾」run 之前，且末尾 run 原文保留。
        """
        from docx import Document
        from lxml import etree

        tool = DocxRevisionTool()
        doc = Document()
        para = doc.add_paragraph()
        prefix_run = para.add_run("前缀AAA")
        para.add_run("中间BBB")
        para.add_run("尾巴CCC")
        tail_run = para.add_run("末尾DDD")

        result = tool.apply_revision(para, "中间BBB尾巴CCC", "替换EE")
        assert result is True

        children = list(para._element)
        prefix_idx = children.index(prefix_run._element)
        tail_idx = children.index(tail_run._element)
        assert prefix_idx < tail_idx

        # del/ins 修订块夹在前缀与末尾之间（旧 bug 会落到段尾，夹层为空）
        in_between = children[prefix_idx + 1 : tail_idx]
        localnames = [etree.QName(c).localname for c in in_between]
        assert "del" in localnames
        assert "ins" in localnames
        # 前缀与末尾 run 原文保留、语序不变
        assert prefix_run.text == "前缀AAA"
        assert tail_run.text == "末尾DDD"
        assert children.index(prefix_run._element) < children.index(tail_run._element)

    def _template_para(self):
        """复刻合同模板「鉴于」段的 run 结构：斜体+下划线的空位 run 夹在普通 run 间。

        真实事故（2026-10-26，/projects/393）：文档里空位是两个空格
        `'与' + '/ 拖欠  '(i+u) + '之事宜'`，编辑器复制出的选区只剩一个空格，
        跨 run 精确 find 必然失配 → 「无法精确定笔」。
        """
        from docx import Document

        doc = Document()
        para = doc.add_paragraph()
        para.add_run("甲方拟")
        para.add_run("就")
        para.add_run("与")
        blank = para.add_run("/ 拖欠  ")
        blank.italic = True
        blank.underline = True
        para.add_run("之事宜")
        para.add_run("，拟委托法律服务机构提供专项法律服务。")
        return para

    def test_apply_revision_whitespace_collapsed_selection(self):
        """选区空白被编辑器折叠（两空格→一空格）也能落笔，del 写文档实际文本。"""
        from docx.oxml.ns import qn

        tool = DocxRevisionTool()
        para = self._template_para()
        assert tool.apply_revision(para, "甲方拟就与/ 拖欠 之事宜", "鉴于甲方拟就对方拖欠货款事宜") is True
        dts = [t.text for t in para._element.iter(qn("w:delText"))]
        # del 内容是文档实际区间（保留双空格），不是选区文本
        assert dts == ["甲方拟就与/ 拖欠  之事宜"]

    def test_apply_revision_whitespace_free_selection(self):
        """选区完全没空格（下划线空位复制时被丢弃）也能落笔。"""
        tool = DocxRevisionTool()
        para = self._template_para()
        assert tool.apply_revision(para, "甲方拟就与/拖欠之事宜", "鉴于甲方拟就对方拖欠货款事宜") is True

    def test_apply_revision_selection_with_extra_space(self):
        """反向差异：选区带空格、文档没有，同样要能落笔。"""
        tool = DocxRevisionTool()
        para = self._template_para()
        assert tool.apply_revision(para, "之事宜， 拟委托", "之事宜，现拟委托") is True

    def test_apply_revision_real_content_difference_still_fails(self):
        """空白容忍不是无底线：真实内容不同必须返回 False（防误删）。"""
        tool = DocxRevisionTool()
        para = self._template_para()
        assert tool.apply_revision(para, "乙方系依法设立的证券公司", "x") is False
        assert "w:del" not in para._element.xml

    def test_find_span_maps_back_to_original_coordinates(self):
        from apps.contract_review.services.formatting.docx_revision_tool import _find_span

        assert _find_span("abc", "abc") == (0, 3)
        assert _find_span("xabcx", "abc") == (1, 4)
        # 去空白等价：区间映射回原坐标（含中间空白、不含尾后空白）
        assert _find_span("A  B C", "A B") == (0, 4)
        assert _find_span("A  B C", "AB C") == (0, 6)
        assert _find_span("A  B C", "zzz") is None
        assert _find_span("A  B C", "   ") is None
