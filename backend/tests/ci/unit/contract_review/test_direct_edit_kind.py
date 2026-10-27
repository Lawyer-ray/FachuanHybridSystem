"""划词直改的合同/案件分流（direct_edit 的 kind 参数化）。

**为什么锁这个**：案件接入前 `_browse` / `_bump_editor_doc_key` / API 层全部
写死合同——case_id 被当 contract_id 发出去会写进**同编号的另一份合同**（静默
写坏，无任何报错）。这些用例守住三件事：

- API 层 `_resolve_target`：二选一硬校验（都给/都不给 422）；
- 服务层按 kind 选 browse service（ACL 随之分流：合同 ContractAccessPolicy、
  案件 CaseAccessPolicy）；
- key 轮换路径按 kind 拼（/contracts/ vs /cases/ + task_type），路径错位
  画布就刷不到新内容。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from apps.contract_review.api.direct_edit_api import _resolve_target
from apps.contract_review.services.formatting.docx_revision_tool import visible_text
from apps.contract_review.services.review import direct_edit_service as mod
from apps.core.exceptions import ValidationException


def _make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestResolveTarget:
    def test_contract_only(self) -> None:
        assert _resolve_target(318, None) == ("contract", 318)

    def test_case_only(self) -> None:
        assert _resolve_target(None, 399) == ("case", 399)

    def test_both_given_rejected(self) -> None:
        with pytest.raises(ValidationException):
            _resolve_target(318, 399)

    def test_neither_given_rejected(self) -> None:
        with pytest.raises(ValidationException):
            _resolve_target(None, None)


class TestBrowseDispatch:
    def test_default_is_contract(self) -> None:
        from apps.contracts.services.contract.integrations.folder_browse_service import ContractFolderBrowseService

        assert isinstance(mod._browse(), ContractFolderBrowseService)

    def test_case_kind_returns_case_service(self) -> None:
        from apps.cases.services.material.case_folder_browse_service import CaseFolderBrowseService

        assert isinstance(mod._browse("case"), CaseFolderBrowseService)


class TestBumpPath:
    @pytest.mark.django_db
    def test_case_kind_bumps_case_folder_path(self, monkeypatch: Any) -> None:
        """案件直改后轮换的是 case_folder 域的 key——路径前缀与 task_type 都要对。"""
        calls: dict[str, Any] = {}

        def fake_bump(source_path: str, task_type: str = "") -> int:
            calls["source_path"] = source_path
            calls["task_type"] = task_type
            return 1

        from apps.onlyoffice.services import callback_service

        monkeypatch.setattr(callback_service, "bump_external_write", fake_bump)

        mod._bump_editor_doc_key(kind="case", owner_id=399, subpath="2-诉状/起诉状.docx")

        assert (
            calls["source_path"]
            == "/api/v1/cases/399/folder-file?subpath=2-%E8%AF%89%E7%8A%B6%2F%E8%B5%B7%E8%AF%89%E7%8A%B6.docx"
        )
        assert calls["task_type"] == "case_folder"

    @pytest.mark.django_db
    def test_contract_kind_keeps_legacy_path(self, monkeypatch: Any) -> None:
        calls: dict[str, Any] = {}

        def fake_bump(source_path: str, task_type: str = "") -> int:
            calls["source_path"] = source_path
            calls["task_type"] = task_type
            return 1

        from apps.onlyoffice.services import callback_service

        monkeypatch.setattr(callback_service, "bump_external_write", fake_bump)

        mod._bump_editor_doc_key(kind="contract", owner_id=318, subpath="1-合同/a.docx")

        assert calls["source_path"] == "/api/v1/contracts/318/folder-file?subpath=1-%E5%90%88%E5%90%8C%2Fa.docx"
        assert calls["task_type"] == "contract_folder"


class TestApplyKindFlow:
    """apply(kind="case") 端到端：案件服务被调用、合同服务绝不被碰。"""

    @pytest.mark.django_db
    def test_case_apply_uses_case_service_and_bumps_case_key(self, monkeypatch: Any, tmp_path: Path) -> None:
        src = _make_docx("请求判令被告支付货款人民币十万元。")
        f = tmp_path / "起诉状.docx"
        f.write_bytes(src)

        saved: dict[str, Any] = {}

        class FakeCaseBrowse:
            def resolve_file(self, **kw: Any) -> Path:
                saved["resolve_owner"] = kw["owner_id"]
                return f

            def save_file_content(self, **kw: Any) -> dict[str, Any]:
                saved.update(kw)
                return {"subpath": kw["subpath"], "name": "起诉状.docx", "size": len(kw["data"])}

        contract_touched: list[str] = []

        class BoomContract:
            def __getattr__(self, name: str) -> Any:
                def _boom(*_a: Any, **_kw: Any) -> Any:
                    contract_touched.append(name)
                    raise AssertionError("案件直改不应触碰合同服务")

                return _boom

        from apps.cases.services.material import case_folder_browse_service as case_mod
        from apps.contracts.services.contract.integrations import folder_browse_service as contract_mod

        monkeypatch.setattr(case_mod, "CaseFolderBrowseService", FakeCaseBrowse)
        monkeypatch.setattr(contract_mod, "ContractFolderBrowseService", BoomContract)

        bumps: list[tuple[str, str]] = []

        def fake_bump(source_path: str, task_type: str = "") -> int:
            bumps.append((source_path, task_type))
            return 1

        from apps.onlyoffice.services import callback_service

        monkeypatch.setattr(callback_service, "bump_external_write", fake_bump)

        result = mod.DirectEditService().apply(
            owner_id=399,
            subpath="2-诉状/起诉状.docx",
            selection="请求判令被告支付货款人民币十万元",
            suggested="请求判令被告立即支付全部货款及逾期利息",
            kind="case",
            user=None,
        )

        assert result["applied"] is True
        assert saved["resolve_owner"] == 399
        # 修订真的写进了产物：改写文本以 <w:ins> 落在段落里（python-docx 原生
        # .text 看不见 w:ins，用仓库的 visible_text 口径断言）
        out_doc = Document(BytesIO(saved["data"]))
        assert "请求判令被告立即支付全部货款及逾期利息" in visible_text(out_doc.paragraphs[0])
        assert "请求判令被告支付货款" not in visible_text(out_doc.paragraphs[0])
        assert bumps == [
            (
                "/api/v1/cases/399/folder-file?subpath=2-%E8%AF%89%E7%8A%B6%2F%E8%B5%B7%E8%AF%89%E7%8A%B6.docx",
                "case_folder",
            )
        ]
        assert contract_touched == []
