"""docx 修订痕迹 → 清单条目的回流测试。

**为什么要回流**：AI 建议进 `ReviewRevision` 表，律师在 OnlyOffice 里自己改的
进 docx 的 `w:ins`/`w:del`。两条路各走各的，左栏清单就看不见律师的手工修改——
他改完一份合同，清单里还挂着「待决策 12 条」，其中几条其实已经改好了。

测重点：
- del+ins 相邻配对成一条「原文→改法」；
- 纯删除 / 纯插入各算一条（它们同样是改动，律师也要确认）；
- 不重复收 AI 自己写的痕迹（否则左栏同一处出现两个条目）；
- 表格里的段落也要收（合同的付款条款常在表格里）。
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.contract_review.services.review.docx_revision_reader import extract_revisions

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _rev(parent: Any, tag: str, text: str, author: str = "张律师") -> None:
    """在段落里插一个 w:del 或 w:ins，内含一个 run。"""
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:id"), "1")
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), "2026-10-25T00:00:00Z")
    run = OxmlElement("w:r")
    t = OxmlElement("w:delText" if tag == "del" else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    el.append(run)
    parent._p.append(el) if hasattr(parent, "_p") else parent.append(el)


def _para(doc: Document, *texts: str) -> Any:
    p = doc.add_paragraph("".join(texts))
    return p


class TestExtractRevisions:
    def test_del_ins_pair_becomes_one_change(self) -> None:
        doc = Document()
        p = doc.add_paragraph("")
        _rev(p, "del", "十日内")
        _rev(p, "ins", "三十日内")
        got = extract_revisions(doc)
        assert len(got) == 1
        assert got[0].original == "十日内"
        assert got[0].suggested == "三十日内"
        assert got[0].author == "张律师"

    def test_delete_only(self) -> None:
        """删了没补：也是一处改动，律师要确认。"""
        doc = Document()
        p = doc.add_paragraph("")
        _rev(p, "del", "本条款不可撤销")
        got = extract_revisions(doc)
        assert len(got) == 1
        assert got[0].original == "本条款不可撤销"
        assert got[0].suggested == ""

    def test_insert_only(self) -> None:
        """凭空加的条款：original 为空，suggested 是新内容。"""
        doc = Document()
        p = doc.add_paragraph("")
        _rev(p, "ins", "本合同自双方盖章之日起生效。")
        got = extract_revisions(doc)
        assert len(got) == 1
        assert got[0].original == ""
        assert got[0].suggested == "本合同自双方盖章之日起生效。"

    def test_no_revisions_returns_empty(self) -> None:
        doc = Document()
        doc.add_paragraph("这是普通正文，没有任何修订。")
        assert extract_revisions(doc) == []

    def test_paragraph_index_recorded(self) -> None:
        """段落索引要能定位，否则左栏「定位」按钮没用。"""
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("第二段")
        p = doc.add_paragraph("")
        _rev(p, "del", "旧")
        _rev(p, "ins", "新")
        got = extract_revisions(doc)
        assert len(got) == 1
        assert got[0].paragraph_index == 2

    def test_multiple_pairs_in_one_paragraph(self) -> None:
        doc = Document()
        p = doc.add_paragraph("")
        _rev(p, "del", "甲")
        _rev(p, "ins", "乙")
        _rev(p, "del", "丙")
        _rev(p, "ins", "丁")
        got = extract_revisions(doc)
        assert [(r.original, r.suggested) for r in got] == [("甲", "乙"), ("丙", "丁")]

    def test_table_paragraphs_are_included(self) -> None:
        """合同的付款/交付条款常在表格里——只看 doc.paragraphs 会漏掉。"""
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell_p = table.rows[0].cells[0].paragraphs[0]
        _rev(cell_p, "del", "月结30天")
        _rev(cell_p, "ins", "月结60天")
        got = extract_revisions(doc)
        assert len(got) == 1
        assert got[0].original == "月结30天"
        assert got[0].suggested == "月结60天"


class TestManualIngest:
    """回流到清单：AI 自己的痕迹不收，人工的收。"""

    @pytest.mark.django_db
    def test_ingest_skips_ai_author(self) -> None:
        from apps.contract_review.models.review_revision import ReviewRevision, RevisionKind
        from apps.contract_review.models.review_task import ReviewTask
        from apps.contract_review.services.review.review_service import _ingest_manual_revisions

        doc = Document()
        p = doc.add_paragraph("")
        _rev(p, "del", "AI删的", author="法穿SI")
        _rev(p, "ins", "AI加的", author="法穿SI")
        p2 = doc.add_paragraph("")
        _rev(p2, "del", "律师删的", author="张律师")
        _rev(p2, "ins", "律师加的", author="张律师")
        buf = io.BytesIO()
        doc.save(buf)

        task = ReviewTask.objects.create(original_file="a.docx", reviewer_name="法穿SI")
        n = _ingest_manual_revisions(task, buf.getvalue())

        assert n == 1, "只收律师的改动，AI 自己写的不能重复收"
        row = ReviewRevision.objects.get(task=task)
        assert row.kind == RevisionKind.MANUAL
        assert row.original == "律师删的"
        assert row.suggested == "律师加的"
        assert "张律师" in row.reason
        assert "改为" in row.title, "标题要说明改了什么，不能只是「律师手改」"

    @pytest.mark.django_db
    def test_manual_kind_not_wiped_by_ai_rerun(self) -> None:
        """AI 重跑刷新自己的 pending，不能把人工条目一起刷掉。"""
        from apps.contract_review.models.review_revision import ReviewRevision, RevisionKind, RevisionStatus
        from apps.contract_review.models.review_task import ReviewTask
        from apps.contract_review.services.review.contract_reviewer import ReviewResult
        from apps.contract_review.services.review.review_service import _persist_revisions

        task = ReviewTask.objects.create(original_file="a.docx", reviewer_name="法穿SI")
        _persist_revisions(
            task, [ReviewResult("手改前", "手改后", "【】律师手改", 1)], kind=RevisionKind.MANUAL
        )
        assert ReviewRevision.objects.filter(task=task, kind=RevisionKind.MANUAL).count() == 1

        # AI 重跑
        _persist_revisions(task, [ReviewResult("AI原文", "AI改法", "【高】问题", 2)])

        assert ReviewRevision.objects.filter(task=task, kind=RevisionKind.MANUAL).count() == 1, (
            "人工条目必须保留"
        )
        ai = ReviewRevision.objects.filter(task=task, kind=RevisionKind.RISK, status=RevisionStatus.PENDING)
        assert ai.count() == 1
        # 序号接着排，不与人工条目撞号
        ordinals = sorted(ReviewRevision.objects.filter(task=task).values_list("ordinal", flat=True))
        assert len(ordinals) == len(set(ordinals)), "序号不能重复"
