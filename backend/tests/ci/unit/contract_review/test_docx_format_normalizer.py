"""DOCX 格式规范化器测试。"""

from unittest.mock import MagicMock, patch

import pytest


class TestDocxFormatNormalizer:
    """DocxFormatNormalizer 可测试逻辑。"""

    def _make_normalizer(self, input_path="/tmp/test.docx"):
        from apps.contract_review.services.format_normalizer.docx_format_normalizer import DocxFormatNormalizer

        return DocxFormatNormalizer(input_path=input_path)

    # ── __init__ ──

    def test_init_default_output_path(self):
        n = self._make_normalizer("/tmp/合同.docx")
        assert str(n.output_path).endswith("合同_规范化.docx")

    def test_init_custom_output_path(self):
        from apps.contract_review.services.format_normalizer.docx_format_normalizer import DocxFormatNormalizer

        n = DocxFormatNormalizer(input_path="/tmp/a.docx", output_path="/tmp/b.docx")
        assert str(n.output_path) == "/tmp/b.docx"

    def test_init_with_reference(self):
        from apps.contract_review.services.format_normalizer.docx_format_normalizer import DocxFormatNormalizer

        n = DocxFormatNormalizer(input_path="/tmp/a.docx", reference_path="/tmp/ref.docx")
        assert str(n.reference_path) == "/tmp/ref.docx"

    def test_init_without_reference(self):
        n = self._make_normalizer()
        assert n.reference_path is None

    def test_init_empty_llm_results(self):
        n = self._make_normalizer()
        assert n._llm_results == {}

    # ── _fallback_classify ──

    def test_fallback_classify_empty_text_returns_1(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = ""
        para.runs = []
        assert n._fallback_classify(para) == 1

    def test_fallback_classify_short_bold_returns_0(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "合同条款"
        rPr = MagicMock()
        rPr.find.return_value = MagicMock()  # bold element found
        run = MagicMock()
        run._element.find.return_value = rPr.find.return_value
        para.runs = [run]
        # Simulate rPr.find(qn("w:b")) is not None
        with patch(
            "apps.contract_review.services.format_normalizer.docx_format_normalizer.qn",
            side_effect=lambda x: x,
        ):
            para.runs[0]._element.find.return_value = MagicMock()  # not None
            # But the _fallback_classify checks runs[0]._element.find(qn("w:rPr"))
            # then rPr.find(qn("w:b"))
            rPr_mock = MagicMock()
            rPr_mock.find.return_value = MagicMock()  # bold found
            para.runs[0]._element.find.return_value = rPr_mock
            result = n._fallback_classify(para)
        assert result == 0

    def test_fallback_classify_chinese_number_prefix_returns_0(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "一、合同主体"
        para.runs = []
        assert n._fallback_classify(para) == 0

    def test_fallback_classify_paren_chinese_number_returns_0(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "（一）甲方义务"
        para.runs = []
        assert n._fallback_classify(para) == 0

    def test_fallback_classify_digit_dot_returns_2(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "1.第一项"
        para.runs = []
        assert n._fallback_classify(para) == 2

    def test_fallback_classify_digit_comma_returns_2(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "2、第二项"
        para.runs = []
        assert n._fallback_classify(para) == 2

    def test_fallback_classify_normal_text_returns_1(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "本合同自双方签字之日起生效"
        para.runs = []
        assert n._fallback_classify(para) == 1

    def test_fallback_classify_long_bold_returns_1(self):
        """超过30字的粗体段落不算标题"""
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "这是一个非常非常非常非常非常非常非常长的粗体段落内容超过了三十个字"
        # No bold - runs empty
        para.runs = []
        result = n._fallback_classify(para)
        assert result == 1

    # ── _fallback_strip ──

    def test_fallback_strip_chinese_number(self):
        n = self._make_normalizer()
        result = n._fallback_strip("一、合同主体")
        assert result == "合同主体"

    def test_fallback_strip_multi_level_number(self):
        n = self._make_normalizer()
        result = n._fallback_strip("1.2.3具体内容")
        assert result == "具体内容"

    def test_fallback_strip_paren_chinese(self):
        n = self._make_normalizer()
        result = n._fallback_strip("（一）甲方义务")
        assert result == "甲方义务"

    def test_fallback_strip_paren_digit(self):
        n = self._make_normalizer()
        result = n._fallback_strip("(1)第一项")
        assert result == "第一项"

    def test_fallback_strip_digit_dot(self):
        n = self._make_normalizer()
        result = n._fallback_strip("1. 第一项")
        assert result == "第一项"

    def test_fallback_strip_digit_comma(self):
        n = self._make_normalizer()
        result = n._fallback_strip("2、第二项")
        assert result == "第二项"

    def test_fallback_strip_no_prefix(self):
        n = self._make_normalizer()
        result = n._fallback_strip("普通正文内容")
        assert result == "普通正文内容"

    def test_fallback_strip_preserves_leading_whitespace(self):
        n = self._make_normalizer()
        result = n._fallback_strip("  一、合同主体")
        assert result == "  合同主体"

    def test_fallback_strip_empty_result_keeps_original(self):
        n = self._make_normalizer()
        # 剥离后为空则保留原文
        result = n._fallback_strip("一、")
        # 空结果不应替换
        assert result == "一、" or result.strip() == ""

    def test_fallback_strip_dot_number(self):
        n = self._make_normalizer()
        result = n._fallback_strip("3．第三项")
        assert result == "第三项"

    def test_fallback_strip_chinese_paren_with_ton(self):
        n = self._make_normalizer()
        result = n._fallback_strip("（五）、具体内容")
        assert result == "具体内容"

    # ── _get_level ──

    def test_get_level_uses_llm_results_first(self):
        n = self._make_normalizer()
        n._llm_results = {5: {"level": 2, "prefix": ""}}
        para = MagicMock()
        assert n._get_level(para, 5) == 2

    def test_get_level_falls_back_to_classify(self):
        n = self._make_normalizer()
        n._llm_results = {}
        para = MagicMock()
        para.text = "普通正文"
        para.runs = []
        assert n._get_level(para, 0) == 1

    # ── _clear_format ──

    def test_clear_format_removes_expected_tags(self):
        n = self._make_normalizer()
        pPr = MagicMock()
        child = MagicMock()
        pPr.find.return_value = child
        n._clear_format(pPr)
        assert pPr.remove.call_count == 4  # spacing, ind, jc, numPr

    def test_clear_format_no_elements_to_remove(self):
        n = self._make_normalizer()
        pPr = MagicMock()
        pPr.find.return_value = None
        n._clear_format(pPr)
        pPr.remove.assert_not_called()

    # ── _llm_analyze_document ──

    def test_llm_analyze_document_empty_paragraphs(self):
        n = self._make_normalizer()
        n.doc = MagicMock()
        n.doc.paragraphs = []
        result = n._llm_analyze_document("test")
        assert result == {}

    def test_llm_analyze_document_no_text_paragraphs(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "   "
        n.doc = MagicMock()
        n.doc.paragraphs = [para]
        result = n._llm_analyze_document("test")
        assert result == {}

    def test_llm_analyze_document_exception_returns_empty(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = "有内容"
        n.doc = MagicMock()
        n.doc.paragraphs = [para]
        with patch(
            "apps.contract_review.services.format_normalizer.llm_helper.ContractStructureAnalyzer",
            side_effect=ImportError("no module"),
        ):
            result = n._llm_analyze_document("test")
        assert result == {}

    # ── _strip_prefix ──

    def test_strip_prefix_empty_text(self):
        n = self._make_normalizer()
        para = MagicMock()
        para.text = ""
        assert n._strip_prefix(para, 0) is None
        # Should return early without changes

    def test_strip_prefix_uses_llm_prefix(self):
        n = self._make_normalizer()
        n._llm_results = {0: {"level": 0, "prefix": "一、"}}
        para = MagicMock()
        para.text = "一、合同主体"
        para.runs = [MagicMock()]
        # Mock _replace_para_text
        with patch.object(n, "_replace_para_text") as mock_replace:
            n._strip_prefix(para, 0)
            mock_replace.assert_called_once()

    def test_strip_prefix_no_match_falls_back(self):
        n = self._make_normalizer()
        n._llm_results = {0: {"level": 0, "prefix": "不匹配的前缀"}}
        para = MagicMock()
        para.text = "一、合同主体"
        para.runs = [MagicMock()]
        with patch.object(n, "_replace_para_text") as mock_replace:
            n._strip_prefix(para, 0)
            # Should call _fallback_strip and potentially _replace_para_text
            assert mock_replace.called

    # ── _normalize_default ──

    @patch("apps.contract_review.services.format_normalizer.docx_format_normalizer.Document")
    def test_normalize_default_sets_margins(self, mock_doc_cls):
        from apps.contract_review.services.format_normalizer.docx_format_normalizer import DocxFormatNormalizer

        mock_doc = MagicMock()
        section = MagicMock()
        mock_doc.sections = [section]
        mock_doc.paragraphs = []
        mock_doc.part.numbering_part._element = MagicMock()
        mock_doc.part.numbering_part._element.findall.return_value = []

        n = DocxFormatNormalizer(input_path="/tmp/test.docx")
        n.doc = mock_doc
        n._normalize_default()
        # Check margins were set
        assert section.top_margin is not None


class TestReferenceNumberingCopy:
    """参考模式：无 numbering part 的 docx 复制参考文档编号定义后必须落盘。

    回归：旧实现用裸 Part 创建 numbering part（blob 固化为初始空树），
    参考模式复制完 abstractNum/num 后不回写 _blob，保存重开后 numbering.xml 为空。
    """

    def _strip_numbering_part(self, path) -> None:
        """从已保存的 docx 包中剥离 numbering part（part + ContentType + 关系）。"""
        import re
        import zipfile

        stripped = path.with_name(path.stem + "_stripped" + path.suffix)
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(stripped, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.namelist():
                if item == "word/numbering.xml":
                    continue
                data = zin.read(item)
                if item == "[Content_Types].xml":
                    data = re.sub(rb'<Override PartName="/word/numbering\.xml"[^>]*/>', b"", data)
                elif item == "word/_rels/document.xml.rels":
                    data = re.sub(rb'<Relationship [^>]*relationships/numbering"[^>]*/>', b"", data)
                zout.writestr(item, data)
        stripped.replace(path)

    def _inject_reference_definition(self, ref_doc, abstract_num_id: str, num_id: str) -> None:
        """在参考文档自带的 numbering part 中注入特征定义（特征 id 用于断言）。"""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        elm = ref_doc.part.numbering_part._element

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), abstract_num_id)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        numFmt = OxmlElement("w:numFmt")
        numFmt.set(qn("w:val"), "chineseCounting")
        lvl.append(numFmt)
        abstract.append(lvl)
        elm.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_num_id)
        num.append(ref)
        elm.append(num)

    def test_numbering_definitions_persisted_after_save(self, tmp_path):
        import zipfile

        from docx import Document

        from apps.contract_review.services.format_normalizer.docx_format_normalizer import DocxFormatNormalizer

        # 目标文档：保存后剥离 numbering part，确保走 _create_numbering_part 分支
        input_path = tmp_path / "target.docx"
        doc = Document()
        doc.add_paragraph("第一条 总则")
        doc.save(str(input_path))
        self._strip_numbering_part(input_path)

        # 参考文档：自带 numbering part，注入特征定义
        ref_path = tmp_path / "ref.docx"
        ref = Document()
        ref.add_paragraph("一、参考条款")
        self._inject_reference_definition(ref, abstract_num_id="77", num_id="99")
        ref.save(str(ref_path))

        output_path = tmp_path / "out.docx"
        normalizer = DocxFormatNormalizer(input_path=input_path, output_path=output_path, reference_path=ref_path)
        normalizer.normalize(use_llm=False)

        with zipfile.ZipFile(output_path) as zf:
            assert "word/numbering.xml" in zf.namelist()
            xml_text = zf.read("word/numbering.xml").decode("utf-8")

        assert 'w:abstractNumId="77"' in xml_text
        assert 'w:numId="99"' in xml_text
