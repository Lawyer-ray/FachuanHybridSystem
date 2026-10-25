"""合同复审业务链路集成测试（真实 DB + 真实 service，仅 mock LLM）。

覆盖用户实际走的那条路：
    上传 docx → 识别当事人 → 确认立场/步骤 → 流水线产出修订 → 落库
    → GET /revisions → POST /resolve（采纳/驳回/撤销）

为什么 mock LLM 而不是 mock service：要验证的是「AI 产出如何变成律师可决策的
清单」以及「决策如何写回」，这些逻辑在 review_service 与 API 层，LLM 是外部依赖。
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from django.test import Client

from apps.contract_review.models.review_revision import ReviewRevision, RevisionSeverity, RevisionStatus
from apps.contract_review.models.review_task import ReviewTask, TaskStatus
from apps.contract_review.services.review.contract_reviewer import ReviewResult
from apps.contract_review.services.review.review_service import _persist_revisions

pytestmark = pytest.mark.django_db


@pytest.fixture
def user(db: Any) -> Any:
    """ReviewTask.user 指向 AUTH_USER_MODEL（organization.Lawyer），不是 auth.User。"""
    from apps.testing.factories import LawyerFactory

    return LawyerFactory(username="review_tester")


@pytest.fixture
def api_client(user: Any) -> Client:
    from ninja_jwt.tokens import RefreshToken

    c = Client()
    c.defaults["HTTP_AUTHORIZATION"] = f"Bearer {RefreshToken.for_user(user).access_token}"
    return c


@pytest.fixture
def task(user: Any) -> ReviewTask:
    return ReviewTask.objects.create(
        user=user,
        original_file="contract_review/uploads/x.docx",
        contract_title="装修工程合同",
        status=TaskStatus.COMPLETED,
    )


REVIEWS = [
    ReviewResult(
        original="甲方于每季度首月十日内支付当季费用。",
        suggested="甲方于收到合规发票后三十日内支付当期费用。",
        reason="【高】付款节点过于提前：钱先付完工期失控。依据：《民法典》第509条。",
        paragraph_index=4,
    ),
    ReviewResult(
        original="任何一方违反本合同约定的，应向守约方承担赔偿责任。",
        suggested="违约方仅就直接损失承担赔偿责任，累计不超过合同总价款的 20%。",
        reason="【中】责任无上限：敞口与合同金额不对等。依据：《民法典》第584条。",
        paragraph_index=8,
    ),
    ReviewResult(
        original="因本合同发生争议，协商不成的，提交仲裁解决。",
        suggested="协商不成的，提交佛山仲裁委员会按其仲裁规则仲裁。",
        reason="【低】仲裁机构未指明：可能致条款无效。依据：《仲裁法》第16条。",
        paragraph_index=12,
    ),
]


class TestReviewPipelineToDecisions:
    """从 AI 产出 → 清单 → 决策的完整闭环"""

    def test_ai_reviews_become_structured_checklist(self, task: ReviewTask) -> None:
        """流水线产出落库后，清单接口能按 ordinal 返回全部条目。"""
        n = _persist_revisions(task, REVIEWS)
        assert n == 3

        rows = list(ReviewRevision.objects.filter(task=task).order_by("ordinal"))
        assert [r.ordinal for r in rows] == [0, 1, 2]
        assert [r.severity for r in rows] == [
            RevisionSeverity.HIGH,
            RevisionSeverity.MEDIUM,
            RevisionSeverity.LOW,
        ]
        # 标题从 reason 里剥掉等级标记并截到首个断句符
        assert rows[0].title == "付款节点过于提前"
        # 原文/改法完整保留（左栏要展示红删线 → 绿底）
        assert rows[0].original.startswith("甲方于每季度首月")
        assert rows[0].suggested.startswith("甲方于收到合规发票")
        assert all(r.status == RevisionStatus.PENDING for r in rows)

    def test_list_endpoint_returns_counts(
        self, api_client: Client, task: ReviewTask
    ) -> None:
        _persist_revisions(task, REVIEWS)
        r = api_client.get(f"/api/v1/contract-review/{task.id}/revisions")
        assert r.status_code == 200
        body = r.json()
        assert body["total"] == 3
        assert body["pending"] == 3
        assert len(body["items"]) == 3
        assert body["items"][0]["severity"] == "high"

    def test_accept_marks_decided_and_decrements_pending(
        self, api_client: Client, task: ReviewTask
    ) -> None:
        _persist_revisions(task, REVIEWS)
        rev = ReviewRevision.objects.filter(task=task, ordinal=0).first()
        assert rev is not None

        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{rev.id}/resolve",
            data={"action": "accept"},
            content_type="application/json",
        )
        assert r.status_code == 200
        body = r.json()
        assert body["status"] == "accepted"
        assert body["pending"] == 2

        rev.refresh_from_db()
        assert rev.status == RevisionStatus.ACCEPTED
        assert rev.decided_at is not None

    def test_reject_keeps_reason_for_ai_feedback(
        self, api_client: Client, task: ReviewTask
    ) -> None:
        """驳回理由可留空；填了要存下来（用于改进后续审查）。"""
        _persist_revisions(task, REVIEWS)
        rev = ReviewRevision.objects.filter(task=task, ordinal=2).first()
        assert rev is not None

        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{rev.id}/resolve",
            data={"action": "reject", "reason": "客户坚持原仲裁条款"},
            content_type="application/json",
        )
        assert r.status_code == 200
        assert r.json()["status"] == "rejected"

        rev.refresh_from_db()
        assert rev.status == RevisionStatus.REJECTED
        assert rev.reject_reason == "客户坚持原仲裁条款"

    def test_reject_without_reason(self, api_client: Client, task: ReviewTask) -> None:
        _persist_revisions(task, REVIEWS)
        rev = ReviewRevision.objects.filter(task=task, ordinal=1).first()
        assert rev is not None
        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{rev.id}/resolve",
            data={"action": "reject"},
            content_type="application/json",
        )
        assert r.status_code == 200
        rev.refresh_from_db()
        assert rev.reject_reason == ""

    def test_reset_undoes_decision(self, api_client: Client, task: ReviewTask) -> None:
        """撤销：回到待决策，且清掉驳回理由与决策时间。"""
        _persist_revisions(task, REVIEWS)
        rev = ReviewRevision.objects.filter(task=task, ordinal=0).first()
        assert rev is not None
        rev.status = RevisionStatus.REJECTED
        rev.reject_reason = "先不改"
        rev.save(update_fields=["status", "reject_reason"])

        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{rev.id}/resolve",
            data={"action": "reset"},
            content_type="application/json",
        )
        assert r.status_code == 200
        assert r.json()["status"] == "pending"

        rev.refresh_from_db()
        assert rev.status == RevisionStatus.PENDING
        assert rev.reject_reason == ""
        assert rev.decided_at is None

    def test_decisions_survive_rerun(self, api_client: Client, task: ReviewTask) -> None:
        """重跑审查不能抹掉律师已做的决策——那是他的判断，不是 AI 的中间产物。"""
        _persist_revisions(task, REVIEWS)
        rev = ReviewRevision.objects.filter(task=task, ordinal=0).first()
        assert rev is not None
        api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{rev.id}/resolve",
            data={"action": "accept"},
            content_type="application/json",
        )

        # 重跑（AI 又给了一批）
        _persist_revisions(task, REVIEWS + [ReviewResult("新增条款", "改法", "【低】新问题。", 20)])

        accepted = ReviewRevision.objects.filter(task=task, status=RevisionStatus.ACCEPTED)
        assert accepted.count() == 1, "已采纳的条目必须保留"
        pending = ReviewRevision.objects.filter(task=task, status=RevisionStatus.PENDING)
        assert pending.count() == 4, "待决策的重跑后应被刷新为最新一批"


class TestAccessControl:
    """别人的任务不能看、不能改"""

    def test_other_user_cannot_list(self, user: Any) -> None:
        from ninja_jwt.tokens import RefreshToken

        from apps.testing.factories import LawyerFactory

        other = LawyerFactory(username="other2")
        t = ReviewTask.objects.create(user=user, original_file="a.docx", status=TaskStatus.COMPLETED)
        _persist_revisions(t, REVIEWS)

        c = Client()
        c.defaults["HTTP_AUTHORIZATION"] = f"Bearer {RefreshToken.for_user(other).access_token}"
        r = c.get(f"/api/v1/contract-review/{t.id}/revisions")
        assert r.status_code == 403

    def test_other_user_cannot_resolve(self, user: Any) -> None:
        from ninja_jwt.tokens import RefreshToken

        from apps.testing.factories import LawyerFactory

        other = LawyerFactory(username="other3")
        t = ReviewTask.objects.create(user=user, original_file="a.docx", status=TaskStatus.COMPLETED)
        _persist_revisions(t, REVIEWS)
        rev = ReviewRevision.objects.filter(task=t).first()
        assert rev is not None

        c = Client()
        c.defaults["HTTP_AUTHORIZATION"] = f"Bearer {RefreshToken.for_user(other).access_token}"
        r = c.post(
            f"/api/v1/contract-review/{t.id}/revisions/{rev.id}/resolve",
            data={"action": "accept"},
            content_type="application/json",
        )
        assert r.status_code == 403
        rev.refresh_from_db()
        assert rev.status == RevisionStatus.PENDING

    def test_missing_revision_is_404(self, api_client: Client, task: ReviewTask) -> None:
        import uuid

        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/revisions/{uuid.uuid4()}/resolve",
            data={"action": "accept"},
            content_type="application/json",
        )
        assert r.status_code == 404


class TestUploadEndpoint:
    """/contract-review/upload 的输入校验"""

    def test_rejects_non_docx(self, api_client: Client) -> None:
        f = io.BytesIO(b"%PDF-1.4 fake")
        f.name = "a.pdf"
        r = api_client.post("/api/v1/contract-review/upload", {"file": f})
        assert r.status_code in (400, 422, 500)
        # 不应创建任务
        assert not ReviewTask.objects.filter(original_file__endswith=".pdf").exists()

    def test_upload_docx_creates_task(
        self, api_client: Client, settings: Any, tmp_path: Any, monkeypatch: Any
    ) -> None:
        """真实 docx 上传 → 建任务（识别当事人依赖 LLM，这里不关心结果，只要任务建出来）。"""
        from docx import Document

        doc = Document()
        doc.add_paragraph("装修工程合同")
        doc.add_paragraph("甲方：某某建材有限公司")
        doc.add_paragraph("乙方：某某装饰工程有限公司")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "contract.docx"

        # 当事人识别依赖 LLM：打桩成固定结果，避免网络
        from apps.contract_review.services.review.party_identifier import PartyIdentifier

        monkeypatch.setattr(
            PartyIdentifier,
            "identify_parties",
            lambda self, paragraphs: {"party_a": "某某建材有限公司", "party_b": "某某装饰工程有限公司"},
            raising=False,
        )

        r = api_client.post("/api/v1/contract-review/upload", {"file": buf})
        assert r.status_code == 200
        body = r.json()
        assert body["task_id"]
        assert ReviewTask.objects.filter(id=body["task_id"]).exists()


class TestClientDemandsPersisted:
    """客户交代必须落到 ReviewTask 并存进 prompt。

    前端曾把 note 变成一句 toast、字段根本没传——律师写了「一定要压工期」，
    模型完全不知道。这条链断过一次，用端到端测试钉住。
    """

    def test_confirm_party_persists_demands_and_stance(
        self, api_client: Client, task: ReviewTask, monkeypatch: Any
    ) -> None:
        # 触发异步流水线会真跑 LLM：打桩掉提交，只验证落库
        from apps.contract_review.services.review import review_service

        monkeypatch.setattr(review_service, "submit_task", lambda *a, **k: None, raising=False)
        import apps.core.tasking as tasking

        monkeypatch.setattr(tasking, "submit_task", lambda *a, **k: None, raising=False)

        task.status = TaskStatus.PARTIES_IDENTIFIED
        task.save(update_fields=["status"])

        r = api_client.post(
            f"/api/v1/contract-review/{task.id}/confirm-party",
            data={
                "represented_party": "party_a",
                "selected_steps": ["contract_review"],
                "client_note": "这单一定要压工期",
                "client_demands": ["这单一定要压工期", "预算不能再加"],
                "our_client_name": "联胜酒店",
                "stance_basis": "合同抬头甲方为联胜酒店",
            },
            content_type="application/json",
        )
        assert r.status_code == 200

        task.refresh_from_db()
        assert task.client_note == "这单一定要压工期"
        assert task.client_demands == ["这单一定要压工期", "预算不能再加"]
        assert task.our_client_name == "联胜酒店"
        assert task.stance_basis == "合同抬头甲方为联胜酒店"
