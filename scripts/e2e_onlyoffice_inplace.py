"""AI 直改「就地落笔 vs 整页重载」真机自检（本地开发环境专用，不入 CI）。

**为什么需要这个脚本**：这套链路（DS 插件桥 → 内核定位 → 回执）有三个部件
（`plugin.js` / 前端 hook / DS 容器），任何一环错位都表现为同一个用户可见
症状——「AI 改完，整页刷一下」。而这个症状**只在真机上出现**，jsdom/单测都
测不到（插件是 DS iframe 里的独立上下文，内核 API 只存在于 callCommand 中）。

判断依据是 **DocEditor 构造次数**：

- `editors` 不变 且 收到 `applied{ok:true}` → **就地落笔，画布未重建**（期望）
- `editors` 增加 → 前端降级换 document key，用户看到整页刷新（异常）

前置：
  - Document Server 已跑： http://127.0.0.1:9980
  - 插件目录已挂载且 **.gz 由 DS 自己生成**（见 deploy/onlyoffice/README.md）
  - 后端依赖可用（读 SystemConfig 里的 JWT_SECRET 签一份临时 config）

用法：
  cd backend && .venv/bin/python ../scripts/e2e_onlyoffice_inplace.py
  cd backend && .venv/bin/python ../scripts/e2e_onlyoffice_inplace.py --headed
"""

from __future__ import annotations

import argparse
import json
import os
import socket
import socketserver
import sys
import tempfile
import threading
from http.server import SimpleHTTPRequestHandler
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
BACKEND = REPO / "backend"
DOCX = BACKEND / "apps" / "onlyoffice" / "tests" / "fixtures" / "_inplace_probe.docx"

DS = "http://127.0.0.1:9980"

# 探针用它自己的 docx：原文来自实际生成的夹具，避免依赖任何真实案卷文件
PROBE_PARAGRAPH = "本条用于校验 AI 就地落笔能否在画布内定位原文。"
PROBE_REPLACEMENT = "本条用于校验 AI 就地落笔已在画布内就地生效。"

FAILS: list[str] = []


def out(*parts: object) -> None:
    sys.stdout.write(" ".join(str(p) for p in parts) + "\n")


def check(name: str, ok: bool, detail: str = "") -> None:
    prefix = "PASS " if ok else "FAIL "
    if not ok:
        FAILS.append(name)
    out(prefix + name + (f"  {detail}" if detail else ""))


def free_port() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = int(s.getsockname()[1])
    s.close()
    return port


def ensure_fixture() -> Path:
    """生成一份带唯一原文的 docx 夹具（重复运行也稳定）。"""
    from docx import Document

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("就地落笔自检夹具（可删）")
    doc.add_paragraph(PROBE_PARAGRAPH)
    doc.add_paragraph("第二段用于确认多层降级不会误伤无关段落。")
    doc.save(str(DOCX))

    # 自校验：夹具必须真的含原文，否则 apply 失败会被误判成插件缺陷
    texts = [p.text for p in Document(str(DOCX)).paragraphs]
    if PROBE_PARAGRAPH not in texts:
        raise RuntimeError(f"夹具生成异常，未包含探针原文: {texts}")
    return DOCX


def get_jwt_secret() -> str:
    """取 ONLYOFFICE_JWT_SECRET。

    优先级：
      1. **容器环境变量**（权威来源）——DS 用的就是这一份，比对才有意义；
      2. SystemConfig（加密值，需 SecretCodec；本机缺 `.env` 里的派生 key 时解不开）。

    为什么先问容器：这个脚本的目的是「验证后端签的 config 能被 DS 接受」，
    密钥取自 DS 自己才不会出现「两边其实不一致但都自洽」的假阳性。
    """
    import shutil
    import subprocess

    docker = shutil.which("docker")
    try:
        if not docker:
            raise RuntimeError("docker 不在 PATH")
        proc = subprocess.run(
            [docker, "exec", "onlyoffice", "printenv", "JWT_SECRET"],
            capture_output=True,
            text=True,
            timeout=15,
        )
        secret = proc.stdout.strip()
        if secret:
            return secret
    except Exception:
        # 取不到 docker secret 属预期内降级（未起容器/无 docker/未注入密钥均会走到这里），
        # 紧接其后有 SystemConfig 回退路径，故无需记录。勿加日志：E2E 脚本输出需保持干净。
        pass

    sys.path.insert(0, str(BACKEND / "apiSystem"))
    sys.path.insert(0, str(BACKEND))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "apiSystem.settings")
    import django

    django.setup()
    from apps.core.models.system_config import SystemConfig
    from apps.core.security.secret_codec import SecretCodec

    obj = SystemConfig.objects.get(key="ONLYOFFICE_JWT_SECRET")
    codec = SecretCodec()
    # 解不开时**不要静默退化**：密钥不对签出来的 config 会被 DS 拒绝，
    # 表现为「文档打开失败」，很难联想到是密钥问题。
    return codec.decrypt(obj.value) if codec.is_encrypted(obj.value) else (obj.value or "")


PAGE = """<!doctype html><html><head><meta charset="utf-8">
<script src="%(ds)s/web-apps/apps/api/documents/api.js"></script></head>
<body style="margin:0"><div id="ph" style="width:1240px;height:900px"></div>
<script>
window.__editors = 0;
window.__pluginWin = null;
window.__applied = null;
window.__log = [];
(function () {
  var iv = setInterval(function () {
    if (window.DocsAPI && window.DocsAPI.DocEditor && !window.DocsAPI.__patched) {
      var orig = window.DocsAPI.DocEditor;
      window.DocsAPI.DocEditor = function () {
        window.__editors += 1;
        return new (Function.prototype.bind.apply(orig, [null].concat([].slice.call(arguments))))();
      };
      window.DocsAPI.__patched = true;
      clearInterval(iv);
    }
  }, 50);
})();
window.addEventListener('message', function (e) {
  var d = e.data;
  if (!d || d.source !== 'fachuan-plugin') return;
  window.__log.push(d.type + (d.result ? ':' + JSON.stringify(d.result) : ''));
  if (d.type === 'ready' || d.type === 'hello') window.__pluginWin = e.source;
  if (d.type === 'applied') window.__applied = d.result;
});
// __send 必须在构造之前定义：DocEditor 抛异常时后面的语句不会执行，
// 否则会得到「__send is not a function」这种掩盖真因的二次错误。
window.__send = function () {
  if (!window.__pluginWin) return 'no-plugin';
  window.__pluginWin.postMessage({source: 'fachuan-page', action: 'apply-revision',
    payload: {orig: %(orig)s, sug: %(sug)s}}, '*');
  return 'sent';
};
try {
  new DocsAPI.DocEditor("ph", %(cfg)s);
} catch (e) {
  window.__bootError = String((e && e.message) || e);
}
</script></body></html>"""


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(self.root_dir), **kwargs)

    root_dir = Path(tempfile.gettempdir())

    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        super().end_headers()

    def do_POST(self):
        self.rfile.read(int(self.headers.get("Content-Length") or 0))
        self.send_header_safe()
        self.wfile.write(json.dumps({"error": 0}).encode())

    def send_header_safe(self):
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()

    def log_message(self, *args):
        pass


def build_config(secret: str, doc_url: str, callback_url: str, doc_key: str) -> dict:
    """与后端 build_editor_config 同口径的最小 config（含插件 autostart）。

    `doc_key` **必须每次运行唯一**：DS 按 document.key 缓存文档内容，
    用固定 key 时第二次运行会直接吃上一次的缓存（上一轮已被改写），
    表现为 apply-revision 报「原文未找到」——这是探针自身的假阴性。
    """
    import jwt

    guid = "asc.{A67B4C8F-1F2A-4B6E-9C3D-5E8F7A2B1D04}"
    cfg = {
        "documentType": "word",
        "document": {
            "fileType": "docx",
            "key": doc_key,
            "title": "probe.docx",
            "url": doc_url,
            "permissions": {
                "edit": True,
                "review": True,
                "comment": True,
                "download": True,
                "print": True,
                "copy": True,
                "chat": False,
            },
        },
        "editorConfig": {
            "mode": "edit",
            "lang": "zh-CN",
            "user": {"id": "1", "name": "probe"},
            "callbackUrl": callback_url,
            "customization": {
                "compactHeader": True,
                "hideRightMenu": True,
                "autosave": True,
                "forcesave": True,
                "uiTheme": "theme-classic-light",
                "review": {"reviewDisplay": "markup", "showReviewChanges": True},
            },
            "plugins": {
                "autostart": [guid],
                "pluginsData": [f"{DS}/sdkjs-plugins/fachuan-rewrite/config.json"],
            },
        },
        "type": "desktop",
        "width": "100%",
        "height": "100%",
    }
    cfg["token"] = jwt.encode(cfg, secret, algorithm="HS256")
    return cfg


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--headed", action="store_true", help="显示浏览器窗口")
    args = ap.parse_args()

    from playwright.sync_api import sync_playwright

    fixture = ensure_fixture()
    secret = get_jwt_secret()
    check("读到 ONLYOFFICE_JWT_SECRET", bool(secret), f"{len(secret)} 字节")

    port = free_port()
    Handler.root_dir = fixture.parent
    srv = socketserver.TCPServer(("0.0.0.0", port), Handler)
    srv.allow_reuse_address = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()

    doc_url = f"http://host.docker.internal:{port}/{fixture.name}"
    cb_url = f"http://host.docker.internal:{port}/callback"
    doc_key = f"inplace-probe-{port}"
    cfg = build_config(secret, doc_url, cb_url, doc_key)

    page_path = fixture.parent / f"_inplace_probe_{port}.html"
    page_path.write_text(
        PAGE
        % {
            "ds": DS,
            "cfg": json.dumps(cfg),
            "orig": json.dumps(PROBE_PARAGRAPH, ensure_ascii=False),
            "sug": json.dumps(PROBE_REPLACEMENT, ensure_ascii=False),
        },
        encoding="utf-8",
    )

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=not args.headed)
        page = browser.new_page(viewport={"width": 1280, "height": 900})
        page_errors: list[str] = []
        page.on("pageerror", lambda e: page_errors.append(str(e)))
        page.on("console", lambda m: page_errors.append(f"{m.type}: {m.text[:200]}") if m.type == "error" else None)
        page.on("requestfailed", lambda r: page_errors.append(f"reqfail {r.url}"))
        page.on("response", lambda r: page_errors.append(f"{r.status} {r.url}") if r.status >= 400 else None)
        page.goto(f"http://127.0.0.1:{port}/{page_path.name}")

        try:
            page.wait_for_function("() => window.__pluginWin !== null", timeout=60000)
            check("插件桥握手（ready）", True)
        except Exception:
            check("插件桥握手（ready）", False, "60s 未收到 ready —— 先查 plugin.js 是否真的下发（.gz 陷阱）")
            browser.close()
            srv.shutdown()
            page_path.unlink(missing_ok=True)
            return 1

        boot_err = page.evaluate("() => window.__bootError || null")
        if boot_err:
            check("DocEditor 构造未抛异常", False, boot_err)
        if page_errors:
            out("      page errors: " + " | ".join(page_errors[:5]))
        before = page.evaluate("() => window.__editors")
        page.evaluate("() => window.__send()")
        try:
            page.wait_for_function("() => window.__applied !== null", timeout=20000)
        except Exception:
            check("收到 applied 回执", False, "20s 超时")
        page.wait_for_timeout(3000)

        applied = page.evaluate("() => window.__applied")
        after = page.evaluate("() => window.__editors")

        check("apply-revision 成功", bool(applied and applied.get("ok")), json.dumps(applied, ensure_ascii=False))
        check("画布未被重建（editors 不增）", after <= before, f"{before} -> {after}")
        check(
            "走的是就地路径（非换 key 重载）",
            bool(applied and applied.get("ok")) and after <= before,
            "VERDICT: IN-PLACE" if after <= before else "VERDICT: RELOADED",
        )
        if applied and applied.get("how"):
            out(f"      定位层级 how={applied['how']}（search/element/normalized/partial 任一都算就地成功）")

        browser.close()

    srv.shutdown()
    page_path.unlink(missing_ok=True)
    # 夹具是每次运行现生成的临时物，不入库（见 .gitignore）
    fixture.unlink(missing_ok=True)

    out("")
    if FAILS:
        out(f"❌ {len(FAILS)} 项失败：" + "；".join(FAILS))
        return 1
    out("✅ 全部通过：AI 直改就地落笔，画布未重建")
    return 0


if __name__ == "__main__":
    sys.exit(main())
